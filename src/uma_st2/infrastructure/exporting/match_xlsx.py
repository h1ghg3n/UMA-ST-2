"""XLSX renderer for the versioned Circle Match Season export projection."""

from __future__ import annotations

from collections import Counter
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
from re import sub
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from uma_st2.application.exporting import (
    MATCH_EXPORT_PROJECTION_VERSION,
    MATCH_EXPORT_WORKBOOK_SCHEMA_VERSION,
    ExportArtifact,
    MatchExportProjection,
)
from uma_st2.shared import normalize_utc_datetime

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_KST = ZoneInfo("Asia/Seoul")
_EXCEL_MAX_ROWS = 1_048_576
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")


class MatchXlsxRenderError(RuntimeError):
    """The complete Match projection cannot be represented as workbook v1."""


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _kst(value: datetime | None) -> str:
    if value is None:
        return ""
    normalized = normalize_utc_datetime(value)
    return normalized.astimezone(_KST).strftime("%Y-%m-%d %H:%M:%S KST")


def _safe_filename(*, season_key: str, generated_at: datetime) -> str:
    normalized = normalize_utc_datetime(generated_at, field_name="generated_at")
    safe_key = sub(r"[^0-9A-Za-z_-]", "_", season_key).strip(" ._")
    stamp = normalized.astimezone(UTC).strftime("%Y%m%dT%H%M%SZ")
    filename = f"match_{safe_key}_{stamp}.xlsx"
    if not safe_key or PurePath(filename).name != filename:
        raise MatchXlsxRenderError("Match export filename is unsafe.")
    return filename


def _append_row(sheet: Worksheet, values: tuple[object, ...] | list[object]) -> None:
    sheet.append([_safe_cell(value) for value in values])


def _style_header(sheet: Worksheet, *, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit_columns(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


class MatchSeasonXlsxRenderer:
    """Render complete Circle Match workbook v1 bytes outside persistence lifetime."""

    def render(
        self,
        projection: MatchExportProjection,
        *,
        generated_at: datetime,
    ) -> ExportArtifact:
        generated_at = normalize_utc_datetime(generated_at, field_name="generated_at")
        if projection.projection_version != MATCH_EXPORT_PROJECTION_VERSION:
            raise MatchXlsxRenderError("Unsupported Match export projection version.")
        self._validate_sheet_bounds(projection)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "시즌 요약"
        matches_sheet = workbook.create_sheet("경기")
        entries_sheet = workbook.create_sheet("출전 및 결과")
        bets_sheet = workbook.create_sheet("베팅 및 정산")
        ratings_sheet = workbook.create_sheet("레이팅")
        workbook.properties.title = f"Circle Match {projection.season.name}"
        workbook.properties.subject = MATCH_EXPORT_WORKBOOK_SCHEMA_VERSION
        workbook.properties.creator = "UMA-ST-2"
        workbook.properties.created = generated_at.replace(tzinfo=None)
        workbook.properties.modified = generated_at.replace(tzinfo=None)

        self._render_summary(summary, projection=projection, generated_at=generated_at)
        self._render_matches(matches_sheet, projection=projection)
        self._render_entries(entries_sheet, projection=projection)
        self._render_bets(bets_sheet, projection=projection)
        self._render_ratings(ratings_sheet, projection=projection)

        buffer = BytesIO()
        try:
            workbook.save(buffer)
            content = buffer.getvalue()
        except Exception as exc:
            raise MatchXlsxRenderError("Match workbook serialization failed.") from exc
        finally:
            workbook.close()
            buffer.close()

        return ExportArtifact(
            filename=_safe_filename(season_key=projection.season.key, generated_at=generated_at),
            media_type=_XLSX_MEDIA_TYPE,
            content=content,
            sha256_hex=sha256(content).hexdigest(),
            generated_at=generated_at,
            source_cutoff=projection.source_cutoff,
            schema_version=MATCH_EXPORT_WORKBOOK_SCHEMA_VERSION,
            projection_version=projection.projection_version,
            scope_type="match_season",
            scope_id=projection.season.key,
            scope_name=projection.season.name,
            row_count=projection.workbook_data_row_count,
        )

    @staticmethod
    def _validate_sheet_bounds(projection: MatchExportProjection) -> None:
        required_rows = (
            1 + len(projection.matches),
            1 + projection.entry_count,
            1 + projection.bet_count,
            1 + projection.rating_transaction_count,
        )
        if any(row_count > _EXCEL_MAX_ROWS for row_count in required_rows):
            raise MatchXlsxRenderError("Complete Match export exceeds one XLSX sheet row limit.")

    @staticmethod
    def _render_summary(
        sheet: Worksheet,
        *,
        projection: MatchExportProjection,
        generated_at: datetime,
    ) -> None:
        _append_row(sheet, ("항목", "값"))
        metadata = (
            ("Season key", projection.season.key),
            ("Season 표시명", projection.season.name),
            ("KST 시작", _kst(projection.season.starts_at)),
            ("KST 종료", _kst(projection.season.ends_at)),
            ("UTC 시작", projection.season.starts_at.isoformat()),
            ("UTC 종료", projection.season.ends_at.isoformat()),
            ("생성 시각", _kst(generated_at)),
            ("Source cutoff", _kst(projection.source_cutoff)),
            ("Projection version", projection.projection_version),
            ("Workbook version", MATCH_EXPORT_WORKBOOK_SCHEMA_VERSION),
            ("Match 수", len(projection.matches)),
            ("Entry 수", projection.entry_count),
            ("Bet 수", projection.bet_count),
            ("Rating transaction 수", projection.rating_transaction_count),
            ("경제 금액 의미", "Persona 합계 receipt는 관련 Bet row에 같은 값으로 표시"),
        )
        for row in metadata:
            _append_row(sheet, row)
        _style_header(sheet, row=1, start_column=1, end_column=2)
        for row in range(2, 2 + len(metadata)):
            sheet.cell(row=row, column=1).fill = _SECTION_FILL
            sheet.cell(row=row, column=1).font = Font(bold=True)

        start = 19
        counters = (
            ("source kind", Counter(match.source_kind.value for match in projection.matches)),
            ("status", Counter(match.status.value for match in projection.matches)),
            ("grade", Counter(match.grade.value for match in projection.matches)),
        )
        for section, counter in counters:
            _append_row(sheet, (section, "Match 수"))
            _style_header(sheet, row=sheet.max_row, start_column=1, end_column=2)
            for key, count in sorted(counter.items()):
                _append_row(sheet, (key, count))
            _append_row(sheet, ("", ""))
        sheet.freeze_panes = f"A{start}"
        _fit_columns(sheet, (30, 72))

    @staticmethod
    def _render_matches(sheet: Worksheet, *, projection: MatchExportProjection) -> None:
        headers = (
            "Match ID",
            "Match 제목",
            "설명",
            "source kind",
            "등급",
            "개최 일정",
            "상태",
            "terminal 사유",
            "경기장",
            "course",
            "surface",
            "거리",
            "방향",
            "layout",
            "condition season",
            "날씨",
            "시간대",
            "주로 상태",
            "완주 시간(ms)",
            "Entry 수",
            "Bet 수",
        )
        _append_row(sheet, headers)
        for match in projection.matches:
            condition = match.condition
            _append_row(
                sheet,
                (
                    match.id,
                    match.name,
                    match.description or "",
                    match.source_kind.value,
                    match.grade.value,
                    _kst(match.scheduled_at),
                    match.status.value,
                    match.terminal_reason or "",
                    match.course.stadium_name,
                    match.course.course_id,
                    match.course.surface.value,
                    match.course.distance,
                    match.course.direction.value,
                    match.course.layout.value,
                    condition.season.value if condition else "",
                    condition.weather.value if condition else "",
                    condition.time_of_day.value if condition else "",
                    condition.track_condition.value if condition else "",
                    match.finish_time_ms or "",
                    len(match.entries),
                    len(match.bets),
                ),
            )
        _finalize_table(
            sheet,
            headers=headers,
            widths=(11, 28, 40, 14, 10, 22, 18, 32, 24, 11, 11, 10, 12, 16, 17, 12, 12, 14, 16, 10, 10),
        )

    @staticmethod
    def _render_entries(sheet: Worksheet, *, projection: MatchExportProjection) -> None:
        headers = (
            "Match ID",
            "Match 제목",
            "source kind",
            "Match 상태",
            "개최 일정",
            "MatchEntry ID",
            "출전 번호",
            "GameAccount ID",
            "GameAccount 표시명",
            "region",
            "owner-at-event 표시명",
            "당시 소속",
            "캐릭터",
            "variant",
            "각질",
            "육성 등급",
            "최종 순위",
            "Rating 처리",
            "인기 순위",
            "착차",
        )
        _append_row(sheet, headers)
        for match in projection.matches:
            for entry in match.entries:
                _append_row(
                    sheet,
                    (
                        match.id,
                        match.name,
                        match.source_kind.value,
                        match.status.value,
                        _kst(match.scheduled_at),
                        entry.id,
                        entry.entry_number,
                        entry.game_account_id,
                        entry.game_account_name,
                        entry.game_region.value,
                        entry.owner_at_event_display_name,
                        entry.affiliation_at_event or "",
                        entry.character_name,
                        entry.variant_name or "",
                        entry.running_style or "",
                        entry.training_grade or "",
                        entry.rank or "",
                        entry.rating_disposition.value if entry.rating_disposition is not None else "",
                        entry.popularity_rank or "",
                        entry.margin or "",
                    ),
                )
        _finalize_table(
            sheet,
            headers=headers,
            widths=(11, 28, 14, 18, 22, 14, 11, 16, 25, 10, 26, 22, 24, 28, 12, 13, 11, 16, 11, 12),
        )

    @staticmethod
    def _render_bets(sheet: Worksheet, *, projection: MatchExportProjection) -> None:
        headers = (
            "Bet ID",
            "Match ID",
            "Match 제목",
            "source kind",
            "Match 상태",
            "개최 일정",
            "표시명",
            "Bet 유형",
            "선택 출전 번호",
            "베팅액",
            "Bet 상태",
            "생성 시각",
            "수정 시각",
            "적용 배당률",
            "Persona 합계 지급액",
            "Persona 합계 지급 회수액",
            "Persona 합계 환불액",
            "환불 사유",
            "환불 시각",
        )
        _append_row(sheet, headers)
        for match in projection.matches:
            entry_numbers = {entry.id: entry.entry_number for entry in match.entries}
            for bet in match.bets:
                selections = " / ".join(str(entry_numbers[entry_id]) for entry_id in bet.selection_entry_ids)
                _append_row(
                    sheet,
                    (
                        bet.id,
                        match.id,
                        match.name,
                        match.source_kind.value,
                        match.status.value,
                        _kst(match.scheduled_at),
                        bet.persona_display_name,
                        bet.bet_type.value,
                        selections,
                        bet.amount,
                        bet.status.value,
                        _kst(bet.created_at),
                        _kst(bet.updated_at),
                        format(bet.applied_odds, ".1f") if bet.applied_odds is not None else "",
                        bet.payout_amount or "",
                        bet.payout_reversal_amount or "",
                        bet.refund_amount or "",
                        bet.refund_reason or "",
                        _kst(bet.refunded_at),
                    ),
                )
        _finalize_table(
            sheet, headers=headers, widths=(11, 11, 28, 14, 18, 22, 24, 12, 22, 11, 12, 22, 22, 14, 21, 24, 21, 32, 22)
        )

    @staticmethod
    def _render_ratings(sheet: Worksheet, *, projection: MatchExportProjection) -> None:
        headers = (
            "RatingTransaction ID",
            "Match ID",
            "Match 제목",
            "source kind",
            "Match 상태",
            "개최 일정",
            "MatchEntry ID",
            "출전 번호",
            "GameAccount ID",
            "GameAccount 표시명",
            "region",
            "owner-at-event 표시명",
            "최종 순위",
            "rule version",
            "변경 전",
            "변량",
            "변경 후",
            "생성 시각",
        )
        _append_row(sheet, headers)
        for match in projection.matches:
            entry_by_id = {entry.id: entry for entry in match.entries}
            for rating in match.ratings:
                entry = entry_by_id[rating.match_entry_id]
                _append_row(
                    sheet,
                    (
                        rating.id,
                        match.id,
                        match.name,
                        match.source_kind.value,
                        match.status.value,
                        _kst(match.scheduled_at),
                        entry.id,
                        entry.entry_number,
                        entry.game_account_id,
                        entry.game_account_name,
                        entry.game_region.value,
                        entry.owner_at_event_display_name,
                        entry.rank or "",
                        rating.rating_rule_version,
                        format(rating.rating_before, ".1f"),
                        format(rating.amount, ".1f"),
                        format(rating.rating_after, ".1f"),
                        _kst(rating.created_at),
                    ),
                )
        _finalize_table(
            sheet, headers=headers, widths=(21, 11, 28, 14, 18, 22, 14, 11, 16, 25, 10, 26, 11, 14, 13, 13, 13, 22)
        )


def _finalize_table(sheet: Worksheet, *, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    _style_header(sheet, row=1, start_column=1, end_column=len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    _fit_columns(sheet, widths)
