"""XLSX renderer for the versioned WIN5 Season export projection."""

from __future__ import annotations

from collections import defaultdict
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
from re import sub
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from uma_st2.application.exporting import (
    WIN5_EXPORT_PROJECTION_VERSION,
    WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION,
    ExportArtifact,
    Win5ExportPick,
    Win5ExportProjection,
    Win5ExportRace,
    Win5ExportResult,
    Win5ExportRound,
)
from uma_st2.domain.win5 import Win5RoundType, Win5SubmissionStatus
from uma_st2.shared import normalize_utc_datetime

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_KST = ZoneInfo("Asia/Seoul")
_EXCEL_MAX_ROWS = 1_048_576
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")


class Win5XlsxRenderError(RuntimeError):
    """The complete projection cannot be represented as workbook v1."""


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _kst(value: datetime | None) -> str:
    if value is None:
        return ""
    normalized = normalize_utc_datetime(value)
    return normalized.astimezone(_KST).strftime("%Y-%m-%d %H:%M:%S KST")


def _safe_filename(*, season_id: int, season_name: str, generated_at: datetime) -> str:
    normalized = normalize_utc_datetime(generated_at, field_name="generated_at")
    title = sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", season_name).strip(" ._")
    title = sub(r"_+", "_", sub(r"\s+", "_", title))[:64].strip(" ._") or f"season-{season_id}"
    stamp = normalized.astimezone(UTC).strftime("%Y%m%dT%H%M%SZ")
    filename = f"win5_{season_id}_{title}_{stamp}.xlsx"
    if PurePath(filename).name != filename:
        raise Win5XlsxRenderError("WIN5 export filename is unsafe.")
    return filename


def _append_row(sheet: Worksheet, values: tuple[object, ...] | list[object]) -> None:
    sheet.append([_safe_cell(value) for value in values])


def _style_header(sheet: Worksheet, *, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit_columns(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _entry_maps(
    projection: Win5ExportProjection,
) -> tuple[
    dict[int, tuple[int, str]],
    dict[int, Win5ExportRace],
    dict[int, Win5ExportRound],
    dict[int, int],
    dict[int, int],
    dict[int, Win5ExportResult],
]:
    entries: dict[int, tuple[int, str]] = {}
    races: dict[int, Win5ExportRace] = {}
    rounds: dict[int, Win5ExportRound] = {}
    race_order: dict[int, int] = {}
    race_round: dict[int, int] = {}
    results: dict[int, Win5ExportResult] = {}
    for round_ in projection.rounds:
        rounds[round_.id] = round_
        for index, race in enumerate(round_.races, start=1):
            races[race.id] = race
            race_order[race.id] = index
            race_round[race.id] = round_.id
            entries.update({entry.id: (entry.gate_number, entry.name) for entry in race.entries})
            results.update({result.id: result for result in race.results})
    return entries, races, rounds, race_order, race_round, results


def _selection_label(
    *,
    round_: Win5ExportRound,
    pick: Win5ExportPick | None,
    entries: dict[int, tuple[int, str]],
) -> str:
    if pick is None:
        return ""
    if round_.round_type == Win5RoundType.NORMAL:
        gate_number, name = entries[pick.race_entry_id]  # type: ignore[index]
        return f"{gate_number} · {name}"
    return str(pick.gate_number)


def _result_label(
    *,
    round_: Win5ExportRound,
    race: Win5ExportRace,
    result: Win5ExportResult | None,
    entries: dict[int, tuple[int, str]],
) -> str:
    if result is None:
        return "void" if race.void_reason is not None else ""
    if round_.round_type == Win5RoundType.NORMAL:
        gate_number, name = entries[result.race_entry_id]  # type: ignore[index]
        return f"{gate_number} · {name}"
    return str(result.gate_number)


class Win5SeasonXlsxRenderer:
    """Render complete workbook v1 bytes without database resources."""

    def render(
        self,
        projection: Win5ExportProjection,
        *,
        generated_at: datetime,
    ) -> ExportArtifact:
        generated_at = normalize_utc_datetime(generated_at, field_name="generated_at")
        if projection.projection_version != WIN5_EXPORT_PROJECTION_VERSION:
            raise Win5XlsxRenderError("Unsupported WIN5 export projection version.")
        self._validate_sheet_bounds(projection)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "시즌 요약"
        rounds_sheet = workbook.create_sheet("라운드")
        judgements_sheet = workbook.create_sheet("제출 및 판정")
        hall_sheet = workbook.create_sheet("명예의 전당")
        workbook.properties.title = f"WIN5 {projection.season_name}"
        workbook.properties.subject = WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION
        workbook.properties.creator = "UMA-ST-2"
        workbook.properties.created = generated_at.replace(tzinfo=None)
        workbook.properties.modified = generated_at.replace(tzinfo=None)

        self._render_summary(summary, projection=projection, generated_at=generated_at)
        self._render_rounds(rounds_sheet, projection=projection)
        self._render_judgements(judgements_sheet, projection=projection)
        self._render_hall(hall_sheet, projection=projection)

        buffer = BytesIO()
        try:
            workbook.save(buffer)
            content = buffer.getvalue()
        except Exception as exc:
            raise Win5XlsxRenderError("WIN5 workbook serialization failed.") from exc
        finally:
            workbook.close()
            buffer.close()

        return ExportArtifact(
            filename=_safe_filename(
                season_id=projection.season_id,
                season_name=projection.season_name,
                generated_at=generated_at,
            ),
            media_type=_XLSX_MEDIA_TYPE,
            content=content,
            sha256_hex=sha256(content).hexdigest(),
            generated_at=generated_at,
            source_cutoff=projection.source_cutoff,
            schema_version=WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION,
            projection_version=projection.projection_version,
            scope_type="win5_season",
            scope_id=str(projection.season_id),
            scope_name=projection.season_name,
            row_count=projection.workbook_data_row_count,
        )

    @staticmethod
    def _validate_sheet_bounds(projection: Win5ExportProjection) -> None:
        required_rows = {
            "시즌 요약": 17 + max(len(projection.season_standings), len(projection.top1_standings)),
            "라운드": 1 + projection.race_count,
            "제출 및 판정": 1 + projection.judgement_row_count,
            "명예의 전당": 1 + len(projection.hall_of_fame_submission_ids),
        }
        if any(row_count > _EXCEL_MAX_ROWS for row_count in required_rows.values()):
            raise Win5XlsxRenderError("Complete WIN5 export exceeds one XLSX sheet row limit.")

    @staticmethod
    def _render_summary(
        sheet: Worksheet,
        *,
        projection: Win5ExportProjection,
        generated_at: datetime,
    ) -> None:
        _append_row(sheet, ("항목", "값"))
        metadata = (
            ("Season ID", projection.season_id),
            ("Season 제목", projection.season_name),
            ("Season 상태", projection.season_status.value),
            ("시작 참고 시각", _kst(projection.starts_at)),
            ("종료 참고 시각", _kst(projection.ends_at)),
            ("생성 시각", _kst(generated_at)),
            ("Source cutoff", _kst(projection.source_cutoff)),
            ("Projection version", projection.projection_version),
            ("Workbook version", WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION),
            ("Round 수", len(projection.rounds)),
            ("Race 수", projection.race_count),
            ("Submission 수", len(projection.submissions)),
            ("참가자 수", projection.participant_count),
        )
        for row in metadata:
            _append_row(sheet, row)
        _style_header(sheet, row=1, start_column=1, end_column=2)
        for row in range(2, 15):
            sheet.cell(row=row, column=1).fill = _SECTION_FILL
            sheet.cell(row=row, column=1).font = Font(bold=True)

        header_row = 17
        for column, value in enumerate(
            ("종합 순위", "표시명", "종합 점수", "", "TOP1 순위", "표시명", "TOP1 점수"),
            start=1,
        ):
            sheet.cell(row=header_row, column=column, value=_safe_cell(value))
        _style_header(sheet, row=header_row, start_column=1, end_column=3)
        _style_header(sheet, row=header_row, start_column=5, end_column=7)

        table_length = max(len(projection.season_standings), len(projection.top1_standings))
        for index in range(table_length):
            season = projection.season_standings[index] if index < len(projection.season_standings) else None
            top1 = projection.top1_standings[index] if index < len(projection.top1_standings) else None
            _append_row(
                sheet,
                (
                    season.rank if season else "",
                    season.display_name if season else "",
                    season.score if season else "",
                    "",
                    top1.rank if top1 else "",
                    top1.display_name if top1 else "",
                    top1.score if top1 else "",
                ),
            )
        if projection.season_standings:
            chart = BarChart()
            chart.title = "종합 점수 상위권"
            chart.y_axis.title = "점수"
            chart.x_axis.title = "표시명"
            chart.height = 7
            chart.width = 13
            upper = min(len(projection.season_standings), 10)
            data = Reference(sheet, min_col=3, min_row=header_row, max_row=header_row + upper)
            categories = Reference(sheet, min_col=2, min_row=header_row + 1, max_row=header_row + upper)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            sheet.add_chart(chart, "I2")
        sheet.freeze_panes = "A18"
        _fit_columns(sheet, (16, 28, 14, 3, 16, 28, 14, 3, 16))

    @staticmethod
    def _render_rounds(sheet: Worksheet, *, projection: Win5ExportProjection) -> None:
        headers = (
            "Round ID",
            "Round 제목",
            "유형",
            "상태",
            "접수 시작",
            "접수 종료",
            "Race 표시순서",
            "Race ID",
            "Race 제목",
            "개최 일정",
            "void 여부",
            "void 사유",
            "void 시각",
            "권위 결과",
            "accepted 제출 수",
            "cancelled 제출 수",
        )
        _append_row(sheet, headers)
        counts: dict[int, dict[Win5SubmissionStatus, int]] = defaultdict(lambda: defaultdict(int))
        for submission in projection.submissions:
            counts[submission.round_id][submission.status] += 1
        entries, _, _, _, _, _ = _entry_maps(projection)
        for round_ in projection.rounds:
            for race_index, race in enumerate(round_.races, start=1):
                result_text = " / ".join(
                    f"{result.position}착: {_result_label(round_=round_, race=race, result=result, entries=entries)}"
                    for result in race.results
                )
                if not result_text and race.void_reason is not None:
                    result_text = "void"
                _append_row(
                    sheet,
                    (
                        round_.id,
                        round_.name,
                        round_.round_type.value,
                        round_.status.value,
                        _kst(round_.opens_at),
                        _kst(round_.closes_at),
                        race_index,
                        race.id,
                        race.name,
                        _kst(race.scheduled_at),
                        "Y" if race.void_reason is not None else "N",
                        race.void_reason or "",
                        _kst(race.voided_at),
                        result_text,
                        counts[round_.id][Win5SubmissionStatus.ACCEPTED],
                        counts[round_.id][Win5SubmissionStatus.CANCELLED],
                    ),
                )
        _style_header(sheet, row=1, start_column=1, end_column=len(headers))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:P{max(1, sheet.max_row)}"
        _fit_columns(sheet, (11, 24, 11, 12, 22, 22, 14, 10, 28, 22, 11, 32, 22, 60, 18, 20))

    @staticmethod
    def _render_judgements(sheet: Worksheet, *, projection: Win5ExportProjection) -> None:
        headers = (
            "Submission ID",
            "Round ID",
            "Round 제목",
            "유형",
            "Round 상태",
            "제출 상태",
            "표시명",
            "Tier",
            "version",
            "생성 시각",
            "수정 시각",
            "Race 표시순서",
            "Race 제목",
            "position",
            "선택",
            "권위 결과",
            "judgement",
            "item 점수",
            "exact 수",
            "wrong_position 수",
            "off_board 수",
            "missing 수",
            "Season 점수 변화",
            "TOP1 점수 변화",
            "Circle Point 보상",
            "판정 시각",
        )
        _append_row(sheet, headers)
        entries, races, rounds, race_order, _, results = _entry_maps(projection)
        for submission in projection.submissions:
            round_ = rounds[submission.round_id]
            event = submission.score_event
            item_by_pick = (
                {item.submission_pick_id: item for item in event.items if item.submission_pick_id is not None}
                if event is not None
                else {}
            )
            rows: list[tuple[int, int, Win5ExportPick | None, object | None]] = []
            for pick in submission.picks:
                rows.append((race_order[pick.race_id], pick.position, pick, item_by_pick.get(pick.id)))
            if event is not None:
                for item in event.items:
                    if item.submission_pick_id is None or item.submission_pick_id not in item_by_pick:
                        rows.append((race_order[item.race_id], item.position, None, item))
            rows.sort(key=lambda row: (row[0], row[1], row[2].id if row[2] is not None else 0))
            for _, position, pick, untyped_item in rows:
                item = untyped_item
                race_id = pick.race_id if pick is not None else item.race_id  # type: ignore[union-attr]
                race = races[race_id]
                matched_result = (
                    results.get(item.matched_result_id)  # type: ignore[union-attr]
                    if item is not None and item.matched_result_id is not None  # type: ignore[union-attr]
                    else next((result for result in race.results if result.position == position), None)
                )
                _append_row(
                    sheet,
                    (
                        submission.id,
                        submission.round_id,
                        round_.name,
                        round_.round_type.value,
                        round_.status.value,
                        submission.status.value,
                        submission.display_name,
                        submission.tier.value,
                        submission.version,
                        _kst(submission.created_at),
                        _kst(submission.updated_at),
                        race_order[race_id],
                        race.name,
                        position,
                        _selection_label(round_=round_, pick=pick, entries=entries),
                        _result_label(round_=round_, race=race, result=matched_result, entries=entries),
                        item.outcome.value if item is not None else "",  # type: ignore[union-attr]
                        item.season_score_delta if item is not None else "",  # type: ignore[union-attr]
                        event.exact_count if event is not None else "",
                        event.wrong_position_count if event is not None else "",
                        event.off_board_count if event is not None else "",
                        event.missing_count if event is not None else "",
                        event.season_score_delta if event is not None else "",
                        event.top1_score_delta if event is not None else "",
                        event.circle_point_reward if event is not None else "",
                        _kst(event.created_at) if event is not None else "",
                    ),
                )
        _style_header(sheet, row=1, start_column=1, end_column=len(headers))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:Z{max(1, sheet.max_row)}"
        _fit_columns(
            sheet,
            (14, 11, 24, 10, 13, 13, 24, 18, 10, 22, 22, 14, 28, 10, 30, 30, 18, 12, 10, 20, 14, 12, 18, 17, 20, 22),
        )
        if sheet.max_row >= 2:
            sheet.conditional_formatting.add(
                f"R2:R{sheet.max_row}",
                ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"),
            )

    @staticmethod
    def _render_hall(sheet: Worksheet, *, projection: Win5ExportProjection) -> None:
        headers = (
            "Submission ID",
            "Round ID",
            "Round 제목",
            "표시명",
            "적중 선택",
            "권위 결과",
            "Season 점수 변화",
            "TOP1 점수 변화",
            "Circle Point 보상",
            "달성 시각",
        )
        _append_row(sheet, headers)
        entries, races, rounds, _, _, _ = _entry_maps(projection)
        submissions = {submission.id: submission for submission in projection.submissions}
        for submission_id in projection.hall_of_fame_submission_ids:
            submission = submissions[submission_id]
            round_ = rounds[submission.round_id]
            event = submission.score_event
            race = round_.races[0]
            picks = sorted(submission.picks, key=lambda pick: (pick.position, pick.id))
            selections = " / ".join(
                f"{pick.position}착: {_selection_label(round_=round_, pick=pick, entries=entries)}" for pick in picks
            )
            result_text = " / ".join(
                f"{result.position}착: "
                f"{_result_label(round_=round_, race=races[result.race_id], result=result, entries=entries)}"
                for result in race.results
            )
            _append_row(
                sheet,
                (
                    submission.id,
                    submission.round_id,
                    round_.name,
                    submission.display_name,
                    selections,
                    result_text,
                    event.season_score_delta,  # type: ignore[union-attr]
                    event.top1_score_delta,  # type: ignore[union-attr]
                    event.circle_point_reward,  # type: ignore[union-attr]
                    _kst(event.created_at),  # type: ignore[union-attr]
                ),
            )
        _style_header(sheet, row=1, start_column=1, end_column=len(headers))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:J{max(1, sheet.max_row)}"
        _fit_columns(sheet, (14, 11, 24, 24, 72, 72, 18, 17, 20, 22))
