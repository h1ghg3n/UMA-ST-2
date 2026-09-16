"""XLSX renderer for the versioned current Circle Point projection."""

from __future__ import annotations

from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from uma_st2.application.exporting import (
    CIRCLE_POINT_EXPORT_PROJECTION_VERSION,
    CIRCLE_POINT_EXPORT_SCOPE_ID,
    CIRCLE_POINT_EXPORT_SCOPE_NAME,
    CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION,
    CirclePointExportProjection,
    ExportArtifact,
)
from uma_st2.shared import normalize_utc_datetime

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_KST = ZoneInfo("Asia/Seoul")
_EXCEL_MAX_ROWS = 1_048_576
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")


class CirclePointXlsxRenderError(RuntimeError):
    """The complete Circle Point projection cannot be represented as workbook v1."""


def _safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _kst(value: datetime) -> str:
    normalized = normalize_utc_datetime(value)
    return normalized.astimezone(_KST).strftime("%Y-%m-%d %H:%M:%S KST")


def _safe_filename(*, generated_at: datetime) -> str:
    normalized = normalize_utc_datetime(generated_at, field_name="generated_at")
    stamp = normalized.astimezone(UTC).strftime("%Y%m%dT%H%M%SZ")
    filename = f"circle-points_{stamp}.xlsx"
    if PurePath(filename).name != filename:
        raise CirclePointXlsxRenderError("Circle Point export filename is unsafe.")
    return filename


def _append_row(sheet: Worksheet, values: tuple[object, ...]) -> None:
    sheet.append([_safe_cell(value) for value in values])


def _style_header(sheet: Worksheet, *, row: int, end_column: int) -> None:
    for column in range(1, end_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _fit_columns(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _finalize_table(sheet: Worksheet, *, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    _style_header(sheet, row=1, end_column=len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    _fit_columns(sheet, widths)


class CirclePointXlsxRenderer:
    """Render complete Circle Point workbook v1 bytes outside persistence lifetime."""

    def render(
        self,
        projection: CirclePointExportProjection,
        *,
        generated_at: datetime,
    ) -> ExportArtifact:
        generated_at = normalize_utc_datetime(generated_at, field_name="generated_at")
        if projection.projection_version != CIRCLE_POINT_EXPORT_PROJECTION_VERSION:
            raise CirclePointXlsxRenderError("Unsupported Circle Point export projection version.")
        self._validate_sheet_bounds(projection)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "요약"
        wallets_sheet = workbook.create_sheet("현재 잔액")
        transactions_sheet = workbook.create_sheet("거래 내역")
        workbook.properties.title = CIRCLE_POINT_EXPORT_SCOPE_NAME
        workbook.properties.subject = CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION
        workbook.properties.creator = "UMA-ST-2"
        workbook.properties.created = generated_at.replace(tzinfo=None)
        workbook.properties.modified = generated_at.replace(tzinfo=None)

        self._render_summary(summary, projection=projection, generated_at=generated_at)
        self._render_wallets(wallets_sheet, projection=projection)
        self._render_transactions(transactions_sheet, projection=projection)

        buffer = BytesIO()
        try:
            workbook.save(buffer)
            content = buffer.getvalue()
        except Exception as exc:
            raise CirclePointXlsxRenderError("Circle Point workbook serialization failed.") from exc
        finally:
            workbook.close()
            buffer.close()

        return ExportArtifact(
            filename=_safe_filename(generated_at=generated_at),
            media_type=_XLSX_MEDIA_TYPE,
            content=content,
            sha256_hex=sha256(content).hexdigest(),
            generated_at=generated_at,
            source_cutoff=projection.source_cutoff,
            schema_version=CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION,
            projection_version=projection.projection_version,
            scope_type="circle_points",
            scope_id=CIRCLE_POINT_EXPORT_SCOPE_ID,
            scope_name=CIRCLE_POINT_EXPORT_SCOPE_NAME,
            row_count=projection.workbook_data_row_count,
        )

    @staticmethod
    def _validate_sheet_bounds(projection: CirclePointExportProjection) -> None:
        if 1 + len(projection.wallets) > _EXCEL_MAX_ROWS or 1 + len(projection.transactions) > _EXCEL_MAX_ROWS:
            raise CirclePointXlsxRenderError("Complete Circle Point export exceeds one XLSX sheet row limit.")

    @staticmethod
    def _render_summary(
        sheet: Worksheet,
        *,
        projection: CirclePointExportProjection,
        generated_at: datetime,
    ) -> None:
        _append_row(sheet, ("항목", "값"))
        metadata = (
            ("생성 시각", _kst(generated_at)),
            ("Source cutoff", _kst(projection.source_cutoff)),
            ("Projection version", projection.projection_version),
            ("Workbook version", CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION),
            ("현재 wallet 수", len(projection.wallets)),
            ("retained transaction 수", len(projection.transactions)),
            ("현재 서클 포인트 합계", projection.current_balance_total),
        )
        for row in metadata:
            _append_row(sheet, row)
        _style_header(sheet, row=1, end_column=2)
        for row in range(2, 2 + len(metadata)):
            sheet.cell(row=row, column=1).fill = _SECTION_FILL
            sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.freeze_panes = "A2"
        _fit_columns(sheet, (30, 72))

    @staticmethod
    def _render_wallets(sheet: Worksheet, *, projection: CirclePointExportProjection) -> None:
        headers = ("Persona ID", "표시명", "Persona 상태", "현재 서클 포인트", "갱신 시각")
        _append_row(sheet, headers)
        for wallet in projection.wallets:
            _append_row(
                sheet,
                (
                    wallet.persona_id,
                    wallet.display_name,
                    wallet.status.value,
                    wallet.balance,
                    _kst(wallet.updated_at),
                ),
            )
        _finalize_table(sheet, headers=headers, widths=(40, 28, 18, 20, 22))

    @staticmethod
    def _render_transactions(sheet: Worksheet, *, projection: CirclePointExportProjection) -> None:
        headers = (
            "PointTransaction ID",
            "Persona ID",
            "표시명",
            "action",
            "변동량",
            "생성 시각",
            "Operation ID",
            "사유",
        )
        _append_row(sheet, headers)
        for transaction in projection.transactions:
            _append_row(
                sheet,
                (
                    transaction.id,
                    transaction.persona_id,
                    transaction.display_name,
                    transaction.action,
                    transaction.amount,
                    _kst(transaction.created_at),
                    transaction.operation_id if transaction.operation_id is not None else "",
                    transaction.reason or "",
                ),
            )
        _finalize_table(sheet, headers=headers, widths=(22, 40, 28, 34, 14, 22, 18, 40))
