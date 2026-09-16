"""Reviewed XLSX adapter for immutable V2 Rating rules."""

from __future__ import annotations

from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from uma_st2.application.rating import RatingRuleSeedError, SeedRatingRuleVersion
from uma_st2.domain.match import MatchGrade
from uma_st2.domain.rating import RatingRule, RatingRuleError

DEFAULT_RATING_RULE_SHEET_NAME = "Rate 기준표"
_GRADE_ALIASES = {
    "GI": MatchGrade.G1,
    "G1": MatchGrade.G1,
    "GII": MatchGrade.G2,
    "G2": MatchGrade.G2,
    "GIII": MatchGrade.G3,
    "G3": MatchGrade.G3,
}


def prepare_rating_rule_workbook(
    path: str | Path,
    *,
    source_identifier: str,
    sheet_name: str = DEFAULT_RATING_RULE_SHEET_NAME,
) -> SeedRatingRuleVersion:
    """Parse one reviewed workbook into the canonical Application seed command."""

    workbook_path = Path(path)
    normalized_sheet_name = sheet_name.strip() if isinstance(sheet_name, str) else ""
    if not normalized_sheet_name:
        raise RatingRuleSeedError("Rating rule sheet name is required.")
    try:
        source_checksum = _sha256_file(workbook_path)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise RatingRuleSeedError("Could not read a valid Rating rule workbook.") from exc
    try:
        if normalized_sheet_name not in workbook.sheetnames:
            raise RatingRuleSeedError(f"Rating rule workbook is missing sheet: {normalized_sheet_name}.")
        rules, source_range = _parse_sheet(workbook[normalized_sheet_name])
    finally:
        workbook.close()
    return SeedRatingRuleVersion(
        source_identifier=source_identifier,
        source_checksum=source_checksum,
        source_sheet_name=normalized_sheet_name,
        source_range=source_range,
        rules=rules,
    )


def _parse_sheet(worksheet: Any) -> tuple[tuple[RatingRule, ...], str]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise RatingRuleSeedError("Rating rule worksheet is empty.")
    parsed: list[RatingRule] = []
    index = 0
    while index < len(rows):
        grade = _read_grade_header(rows[index])
        if grade is None:
            index += 1
            continue
        if index + 2 >= len(rows):
            raise RatingRuleSeedError(f"Rating rule grade block {grade.value} is incomplete.")
        participant_counts = _read_participant_counts(rows[index + 1], grade=grade)
        index += 2
        while index < len(rows) and _read_grade_header(rows[index]) is None:
            data_row = rows[index]
            converted_rank = _positive_integer(_cell(data_row, 1))
            if converted_rank is not None:
                for column_index, participant_count in participant_counts.items():
                    value = _cell(data_row, column_index)
                    if value is None:
                        continue
                    try:
                        parsed.append(
                            RatingRule(
                                grade=grade,
                                participant_count=participant_count,
                                converted_rank=converted_rank,
                                base_delta=_decimal(value),
                            )
                        )
                    except (RatingRuleError, ValueError) as error:
                        raise RatingRuleSeedError(str(error)) from error
            index += 1
    if not parsed:
        raise RatingRuleSeedError("Rating rule worksheet contains no canonical rules.")
    max_row, max_column = _used_bounds(rows)
    return tuple(parsed), f"A1:{_column_name(max_column)}{max_row}"


def _read_grade_header(row: tuple[Any, ...]) -> MatchGrade | None:
    value = _cell(row, 0)
    if not isinstance(value, str) or _cell(row, 2) != "레이스 인원":
        return None
    normalized = value.strip().upper()
    try:
        return _GRADE_ALIASES[normalized]
    except KeyError as error:
        raise RatingRuleSeedError(f"Unsupported Rating workbook grade: {normalized}.") from error


def _read_participant_counts(row: tuple[Any, ...], *, grade: MatchGrade) -> dict[int, int]:
    counts = {
        column_index: participant_count
        for column_index in range(2, len(row))
        if (participant_count := _positive_integer(row[column_index])) is not None
    }
    if not counts:
        raise RatingRuleSeedError(f"Rating rule grade block {grade.value} has no participant counts.")
    return counts


def _cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _positive_integer(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value)
    return None


def _decimal(value: object) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise RatingRuleSeedError("Rating rule base delta must be numeric.")
    return Decimal(str(value))


def _used_bounds(rows: list[tuple[Any, ...]]) -> tuple[int, int]:
    max_row = 0
    max_column = 0
    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            if value is not None:
                max_row = max(max_row, row_number)
                max_column = max(max_column, column_number)
    return max_row, max_column


def _column_name(column_number: int) -> str:
    letters = ""
    remaining = column_number
    while remaining:
        remaining, offset = divmod(remaining - 1, 26)
        letters = chr(ord("A") + offset) + letters
    return letters


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
