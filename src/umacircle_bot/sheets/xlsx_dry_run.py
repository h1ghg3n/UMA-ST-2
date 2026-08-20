from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class WorkbookKind(StrEnum):
    ROOM_MATCH = "room_match"
    WIN5 = "win5"
    UNKNOWN = "unknown"


REQUIRED_SHEETS_BY_KIND = {
    WorkbookKind.ROOM_MATCH: ("베팅 Data", "룸매치 Data", "동친 포인트", "플레이어"),
    WorkbookKind.WIN5: ("스코어",),
}

OPTIONAL_SHEETS_BY_KIND = {
    WorkbookKind.ROOM_MATCH: ("정답배당", "룸매치 기록", "룸매치 계획표", "Data"),
    WorkbookKind.WIN5: ("Win5 일본", "Win5 국내", "Win5 특별 라운드", "명예의 전당"),
}

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass(frozen=True)
class DryRunIssue:
    severity: str
    code: str
    message: str
    sheet_name: str | None = None


@dataclass(frozen=True)
class SheetDryRunReport:
    name: str
    max_row: int
    max_column: int
    sampled_rows: tuple[tuple[Any, ...], ...]
    non_empty_sampled_rows: int


@dataclass(frozen=True)
class WorkbookDryRunReport:
    path: Path
    kind: WorkbookKind
    sheet_names: tuple[str, ...]
    sheets: tuple[SheetDryRunReport, ...]
    issues: tuple[DryRunIssue, ...]

    @property
    def has_errors(self) -> bool:
        return any(issue.severity == "error" for issue in self.issues)


def inspect_xlsx_workbook(
    path: str | Path,
    *,
    kind: WorkbookKind | str = WorkbookKind.UNKNOWN,
    sheets: Sequence[str] | None = None,
    sample_rows: int = 5,
    sample_columns: int = 12,
) -> WorkbookDryRunReport:
    workbook_path = normalize_xlsx_workbook_path(path)
    workbook_kind = _normalize_workbook_kind(kind)
    _validate_sample_limits(sample_rows=sample_rows, sample_columns=sample_columns)

    workbook = load_workbook(
        filename=workbook_path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        sheet_names = tuple(workbook.sheetnames)
        target_sheet_names = _resolve_target_sheet_names(
            available_sheet_names=sheet_names,
            requested_sheet_names=sheets,
            workbook_kind=workbook_kind,
        )

        reports = tuple(
            _inspect_worksheet(
                workbook[sheet_name],
                sample_rows=sample_rows,
                sample_columns=sample_columns,
            )
            for sheet_name in target_sheet_names
        )
        issues = tuple(
            _collect_issues(sheet_names=sheet_names, requested_sheet_names=sheets, workbook_kind=workbook_kind)
        )
        return WorkbookDryRunReport(
            path=workbook_path,
            kind=workbook_kind,
            sheet_names=sheet_names,
            sheets=reports,
            issues=issues,
        )
    finally:
        workbook.close()


def _inspect_worksheet(worksheet: Any, *, sample_rows: int, sample_columns: int) -> SheetDryRunReport:
    declared_max_row = worksheet.max_row or 0
    declared_max_column = worksheet.max_column or 0
    iter_max_row = min(declared_max_row or sample_rows, sample_rows)
    iter_max_column = min(declared_max_column or sample_columns, sample_columns)

    sampled_rows: list[tuple[Any, ...]] = []
    non_empty_sampled_rows = 0
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=iter_max_row,
        min_col=1,
        max_col=iter_max_column,
        values_only=True,
    ):
        normalized_row = tuple(_normalize_cell_value(value) for value in row)
        if any(value is not None for value in normalized_row):
            non_empty_sampled_rows += 1
        sampled_rows.append(normalized_row)

    return SheetDryRunReport(
        name=worksheet.title,
        max_row=declared_max_row,
        max_column=declared_max_column,
        sampled_rows=tuple(sampled_rows),
        non_empty_sampled_rows=non_empty_sampled_rows,
    )


def _collect_issues(
    *,
    sheet_names: Sequence[str],
    requested_sheet_names: Sequence[str] | None,
    workbook_kind: WorkbookKind,
) -> Iterable[DryRunIssue]:
    available = set(sheet_names)
    requested = set(requested_sheet_names or ())
    for sheet_name in sorted(requested - available):
        yield DryRunIssue(
            severity="error",
            code="requested_sheet_missing",
            message=f"requested sheet is missing: {sheet_name}",
            sheet_name=sheet_name,
        )

    for sheet_name in REQUIRED_SHEETS_BY_KIND.get(workbook_kind, ()):
        if sheet_name not in available:
            yield DryRunIssue(
                severity="error",
                code="required_sheet_missing",
                message=f"required {workbook_kind.value} sheet is missing: {sheet_name}",
                sheet_name=sheet_name,
            )

    for sheet_name in OPTIONAL_SHEETS_BY_KIND.get(workbook_kind, ()):
        if sheet_name not in available:
            yield DryRunIssue(
                severity="warning",
                code="optional_sheet_missing",
                message=f"optional {workbook_kind.value} sheet is missing: {sheet_name}",
                sheet_name=sheet_name,
            )


def _resolve_target_sheet_names(
    *,
    available_sheet_names: Sequence[str],
    requested_sheet_names: Sequence[str] | None,
    workbook_kind: WorkbookKind,
) -> tuple[str, ...]:
    if requested_sheet_names is not None:
        requested = tuple(dict.fromkeys(_normalize_sheet_name(sheet_name) for sheet_name in requested_sheet_names))
        return tuple(sheet_name for sheet_name in requested if sheet_name in available_sheet_names)

    known_sheets = REQUIRED_SHEETS_BY_KIND.get(workbook_kind, ()) + OPTIONAL_SHEETS_BY_KIND.get(workbook_kind, ())
    if known_sheets:
        return tuple(sheet_name for sheet_name in known_sheets if sheet_name in available_sheet_names)
    return tuple(available_sheet_names)


def normalize_xlsx_workbook_path(path: str | Path) -> Path:
    workbook_path = Path(path).expanduser()
    if not workbook_path.exists():
        raise FileNotFoundError(f"workbook not found: {workbook_path}")
    if not workbook_path.is_file():
        raise ValueError(f"workbook path is not a file: {workbook_path}")
    if workbook_path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError("workbook must be an .xlsx or .xlsm file")
    return workbook_path.resolve()


def _normalize_workbook_kind(value: WorkbookKind | str) -> WorkbookKind:
    if isinstance(value, WorkbookKind):
        return value
    if not isinstance(value, str):
        raise ValueError("workbook kind must be text")
    try:
        return WorkbookKind(value.strip().lower())
    except ValueError as exc:
        raise ValueError("unsupported workbook kind") from exc


def _normalize_sheet_name(value: str) -> str:
    if not isinstance(value, str):
        raise ValueError("sheet name must be text")
    normalized = value.strip()
    if not normalized:
        raise ValueError("sheet name must not be blank")
    return normalized


def _validate_sample_limits(*, sample_rows: int, sample_columns: int) -> None:
    if not isinstance(sample_rows, int) or isinstance(sample_rows, bool) or not 1 <= sample_rows <= 50:
        raise ValueError("sample_rows must be an integer between 1 and 50")
    if not isinstance(sample_columns, int) or isinstance(sample_columns, bool) or not 1 <= sample_columns <= 100:
        raise ValueError("sample_columns must be an integer between 1 and 100")


def _normalize_cell_value(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped or None
    return value
