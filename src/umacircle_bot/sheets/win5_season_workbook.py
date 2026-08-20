from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from umacircle_bot.domain.win5 import Win5RoundType, Win5SeasonStatus, Win5SubmissionStatus
from umacircle_bot.services.win5_export_types import Win5SeasonSnapshot

WIN5_EXPORT_SHEETS = ("시즌 요약", "라운드", "제출 및 판정", "명예의 전당")
KST = timezone(timedelta(hours=9), name="KST")

_GREEN = "FF49B73F"
_DARK_GREEN = "FF287D2D"
_LIGHT_GREEN = "FFE8F5E9"
_PALE_GREEN = "FFF3FAF3"
_GOLD = "FFFFC928"
_LIGHT_GOLD = "FFFFF7D1"
_RED = "FFFDECEC"
_GRAY = "FFF3F4F6"
_DARK_TEXT = "FF1F2937"
_WHITE = "FFFFFFFF"
_THIN_GRAY = Side(style="thin", color="FFD1D5DB")
_TABLE_BORDER = Border(left=_THIN_GRAY, right=_THIN_GRAY, top=_THIN_GRAY, bottom=_THIN_GRAY)

_ROUND_TYPE_LABELS = {
    Win5RoundType.NORMAL.value: "일반",
    Win5RoundType.SPECIAL.value: "특별",
}
_ROUND_STATUS_LABELS = {
    "setup": "준비",
    "open": "접수 중",
    "closed": "접수 마감",
    "result_entered": "결과 입력",
    "scored": "채점 완료",
}
_SEASON_STATUS_LABELS = {
    Win5SeasonStatus.ACTIVE.value: "진행 중",
    Win5SeasonStatus.CLOSED.value: "종료",
}
_TIER_LABELS = {
    "top1": "TOP1",
    "top3": "TOP3",
    "top5": "TOP5",
    "special_winner": "특별 1착",
}
_SUBMISSION_STATUS_LABELS = {
    Win5SubmissionStatus.ACCEPTED.value: "제출",
    Win5SubmissionStatus.CANCELLED.value: "취소",
}


def write_win5_season_workbook(output_path: Path, *, snapshot: Win5SeasonSnapshot) -> None:
    """Write one season snapshot without formulas or domain recalculation."""

    workbook = Workbook()
    try:
        summary = workbook.active
        summary.title = WIN5_EXPORT_SHEETS[0]
        rounds = workbook.create_sheet(WIN5_EXPORT_SHEETS[1])
        submissions = workbook.create_sheet(WIN5_EXPORT_SHEETS[2])
        hall_of_fame = workbook.create_sheet(WIN5_EXPORT_SHEETS[3])
        _write_summary_sheet(summary, snapshot=snapshot)
        _write_round_sheet(rounds, snapshot=snapshot)
        _write_submission_sheet(submissions, snapshot=snapshot)
        _write_hall_of_fame_sheet(hall_of_fame, snapshot=snapshot)

        with NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as temporary:
            temporary_path = Path(temporary.name)
        try:
            workbook.save(temporary_path)
            temporary_path.replace(output_path)
        except Exception:
            temporary_path.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()


def _write_summary_sheet(worksheet: Worksheet, *, snapshot: Win5SeasonSnapshot) -> None:
    _sheet_base(worksheet, tab_color=_GREEN)
    _title_block(
        worksheet,
        title=f"WIN5 {snapshot.season_number}시즌 · {snapshot.season_name}",
        subtitle=f"시즌 단위 snapshot · 생성 {_format_datetime(snapshot.generated_at)}",
        end_column=8,
    )
    cards = (
        ("상태", _SEASON_STATUS_LABELS[snapshot.season_status]),
        ("라운드", snapshot.round_count),
        ("참가자", snapshot.participant_count),
        ("채점 완료", snapshot.scored_round_count),
        ("시작", _format_datetime(snapshot.starts_at, date_only=True)),
        ("종료", _format_datetime(snapshot.ends_at, date_only=True)),
        ("제출 내역", len(snapshot.submissions)),
        ("TOP5 완전 적중", len(snapshot.hall_of_fame)),
    )
    for index, (label, value) in enumerate(cards):
        row = 5 + index // 4 * 2
        column = 1 + index % 4 * 2
        label_cell = worksheet.cell(row=row, column=column, value=label)
        value_cell = worksheet.cell(row=row, column=column + 1, value=_safe_cell(value))
        label_cell.fill = PatternFill("solid", fgColor=_DARK_GREEN)
        label_cell.font = Font(name="Malgun Gothic", bold=True, color=_WHITE)
        value_cell.fill = PatternFill("solid", fgColor=_PALE_GREEN)
        value_cell.font = Font(name="Malgun Gothic", bold=True, color=_DARK_TEXT)
        for cell in (label_cell, value_cell):
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = _TABLE_BORDER

    headers = (
        "종합 순위",
        "참가자",
        "시즌 점수",
        "TOP1 순위",
        "TOP1 점수",
        "참가 라운드",
        "완전 적중",
        "최근 판정",
    )
    header_row = 10
    _append_headers(worksheet, row=header_row, headers=headers)
    for standing in snapshot.standings:
        worksheet.append(
            (
                standing.season_rank,
                _safe_cell(standing.participant_name),
                standing.season_score,
                standing.top1_rank,
                standing.top1_score,
                standing.participated_round_count,
                standing.perfect_top5_count,
                _format_datetime(standing.latest_judged_at),
            )
        )
        _style_data_row(worksheet, row=worksheet.max_row, columns=len(headers))
    _finish_table(
        worksheet,
        header_row=header_row,
        column_count=len(headers),
        table_name="Win5SeasonStandings",
    )
    if snapshot.standings:
        worksheet.conditional_formatting.add(
            f"C{header_row + 1}:C{worksheet.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="FFF3FAF3",
                end_type="max",
                end_color="FF49B73F",
            ),
        )
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "시즌 점수 TOP 10"
        chart.height = 7
        chart.width = 12
        top_count = min(len(snapshot.standings), 10)
        chart.add_data(
            Reference(
                worksheet,
                min_col=3,
                min_row=header_row,
                max_row=header_row + top_count,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                worksheet,
                min_col=2,
                min_row=header_row + 1,
                max_row=header_row + top_count,
            )
        )
        chart.legend = None
        worksheet.add_chart(chart, "J2")
    worksheet.freeze_panes = f"A{header_row + 1}"
    _set_widths(worksheet, (12, 20, 12, 12, 12, 14, 14, 22))


def _write_round_sheet(worksheet: Worksheet, *, snapshot: Win5SeasonSnapshot) -> None:
    _sheet_base(worksheet, tab_color=_DARK_GREEN)
    headers = (
        "라운드",
        "라운드명",
        "구분",
        "상태",
        "경기",
        "경기 시작",
        "결과",
        "제출",
        "취소",
        "오픈",
        "마감",
        "생성",
        "수정",
    )
    _title_block(
        worksheet,
        title=f"{snapshot.season_name} · 라운드 및 결과",
        subtitle="특별 라운드는 경기별 행으로 표시합니다.",
        end_column=len(headers),
    )
    header_row = 5
    _append_headers(worksheet, row=header_row, headers=headers)
    for row in snapshot.rounds:
        round_label = row.round_label or ""
        if row.race_display_order is not None:
            round_label = (
                f"{round_label} · {row.race_display_order}경기" if round_label else f"{row.race_display_order}경기"
            )
        worksheet.append(
            (
                f"{row.round_number}R",
                _safe_cell(round_label),
                _ROUND_TYPE_LABELS.get(row.round_type, row.round_type),
                _ROUND_STATUS_LABELS.get(row.round_status, row.round_status),
                _safe_cell(row.race_name or ""),
                _format_datetime(row.race_starts_at),
                _safe_cell(" / ".join(row.result_labels)),
                row.accepted_submission_count,
                row.cancelled_submission_count,
                _format_datetime(row.opens_at),
                _format_datetime(row.closes_at),
                _format_datetime(row.created_at),
                _format_datetime(row.updated_at),
            )
        )
        _style_data_row(worksheet, row=worksheet.max_row, columns=len(headers))
    _finish_table(
        worksheet,
        header_row=header_row,
        column_count=len(headers),
        table_name="Win5SeasonRounds",
    )
    worksheet.freeze_panes = f"A{header_row + 1}"
    _set_widths(worksheet, (10, 22, 10, 12, 24, 20, 50, 10, 10, 20, 20, 20, 20))


def _write_submission_sheet(worksheet: Worksheet, *, snapshot: Win5SeasonSnapshot) -> None:
    _sheet_base(worksheet, tab_color=_GREEN)
    headers = (
        "라운드",
        "라운드명",
        "구분",
        "경기",
        "참가자",
        "티어",
        "제출 상태",
        "선택 1",
        "선택 2",
        "선택 3",
        "선택 4",
        "선택 5",
        "결과 1",
        "결과 2",
        "결과 3",
        "결과 4",
        "결과 5",
        "정확 적중",
        "순위권 적중",
        "미적중",
        "시즌 점수",
        "TOP1 점수",
        "서클 포인트",
        "제출 시각",
        "판정 시각",
    )
    _title_block(
        worksheet,
        title=f"{snapshot.season_name} · 제출 및 판정",
        subtitle="모든 값은 export 시점의 persisted snapshot이며 workbook에서 다시 계산하지 않습니다.",
        end_column=len(headers),
    )
    header_row = 5
    _append_headers(worksheet, row=header_row, headers=headers)
    for submission in snapshot.submissions:
        picks = _pad(submission.pick_labels, 5)
        results = _pad(submission.result_labels, 5)
        worksheet.append(
            (
                f"{submission.round_number}R",
                _safe_cell(submission.round_label or ""),
                _ROUND_TYPE_LABELS.get(submission.round_type, submission.round_type),
                _safe_cell(submission.race_name or ""),
                _safe_cell(submission.participant_name),
                _TIER_LABELS.get(submission.prediction_tier, submission.prediction_tier),
                _SUBMISSION_STATUS_LABELS.get(submission.status, submission.status),
                *(_safe_cell(value) for value in picks),
                *(_safe_cell(value) for value in results),
                submission.exact_position_count,
                submission.on_board_wrong_position_count,
                submission.off_board_count,
                submission.season_score_delta,
                submission.top1_score_delta,
                submission.circle_point_reward,
                _format_datetime(submission.submitted_at),
                _format_datetime(submission.judged_at),
            )
        )
        row_number = worksheet.max_row
        _style_data_row(worksheet, row=row_number, columns=len(headers))
        if submission.status == Win5SubmissionStatus.CANCELLED.value:
            for cell in worksheet[row_number]:
                cell.fill = PatternFill("solid", fgColor=_GRAY)
        elif submission.judged_at is not None:
            for index, pick_number in enumerate(submission.pick_numbers[:5]):
                cell = worksheet.cell(row=row_number, column=8 + index)
                if index < len(submission.result_numbers) and pick_number == submission.result_numbers[index]:
                    cell.fill = PatternFill("solid", fgColor=_LIGHT_GREEN)
                elif pick_number in submission.result_numbers:
                    cell.fill = PatternFill("solid", fgColor=_LIGHT_GOLD)
                else:
                    cell.fill = PatternFill("solid", fgColor=_RED)
    _finish_table(
        worksheet,
        header_row=header_row,
        column_count=len(headers),
        table_name="Win5SeasonSubmissions",
    )
    worksheet.freeze_panes = f"H{header_row + 1}"
    _set_widths(
        worksheet,
        (
            10,
            20,
            10,
            24,
            20,
            12,
            12,
            22,
            22,
            22,
            22,
            22,
            22,
            22,
            22,
            22,
            22,
            12,
            14,
            10,
            12,
            12,
            14,
            20,
            20,
        ),
    )


def _write_hall_of_fame_sheet(worksheet: Worksheet, *, snapshot: Win5SeasonSnapshot) -> None:
    _sheet_base(worksheet, tab_color=_GOLD)
    headers = ("기록", "참가자", "라운드", "경기", "1착", "2착", "3착", "4착", "5착", "판정 시각")
    _title_block(
        worksheet,
        title="WIN5 TOP5 완전 적중 명예의 전당",
        subtitle=f"{snapshot.season_number}시즌 · 다섯 위치를 모두 정확히 맞힌 persisted 기록",
        end_column=len(headers),
    )
    header_row = 5
    _append_headers(
        worksheet,
        row=header_row,
        headers=headers,
        fill_color=_GOLD,
        font_color=_DARK_TEXT,
    )
    for index, submission in enumerate(snapshot.hall_of_fame, start=1):
        results = _pad(submission.result_labels, 5)
        worksheet.append(
            (
                f"시즌 {index}호",
                _safe_cell(submission.participant_name),
                f"{submission.round_number}R",
                _safe_cell(submission.race_name or ""),
                *(_safe_cell(value) for value in results),
                _format_datetime(submission.judged_at),
            )
        )
        _style_data_row(
            worksheet,
            row=worksheet.max_row,
            columns=len(headers),
            fill_color=_LIGHT_GOLD,
        )
    if snapshot.hall_of_fame:
        _finish_table(
            worksheet,
            header_row=header_row,
            column_count=len(headers),
            table_name="Win5SeasonHallOfFame",
        )
    else:
        worksheet.merge_cells(
            start_row=header_row + 1,
            start_column=1,
            end_row=header_row + 1,
            end_column=len(headers),
        )
        cell = worksheet.cell(row=header_row + 1, column=1, value="아직 TOP5 완전 적중 기록이 없습니다.")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill("solid", fgColor=_LIGHT_GOLD)
        cell.font = Font(name="Malgun Gothic", italic=True, color=_DARK_TEXT)
    worksheet.freeze_panes = f"A{header_row + 1}"
    _set_widths(worksheet, (14, 20, 10, 24, 22, 22, 22, 22, 22, 20))


def _sheet_base(worksheet: Worksheet, *, tab_color: str) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = tab_color


def _title_block(worksheet: Worksheet, *, title: str, subtitle: str, end_column: int) -> None:
    worksheet.merge_cells(start_row=1, start_column=1, end_row=2, end_column=end_column)
    title_cell = worksheet.cell(row=1, column=1, value=_safe_cell(title))
    title_cell.fill = PatternFill("solid", fgColor=_GREEN)
    title_cell.font = Font(name="Malgun Gothic", size=20, bold=True, color=_WHITE)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 27
    worksheet.row_dimensions[2].height = 14
    for row in (1, 2):
        for column in range(1, end_column + 1):
            worksheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=_GREEN)
    worksheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=end_column)
    subtitle_cell = worksheet.cell(row=3, column=1, value=_safe_cell(subtitle))
    subtitle_cell.fill = PatternFill("solid", fgColor=_LIGHT_GREEN)
    subtitle_cell.font = Font(name="Malgun Gothic", size=10, color=_DARK_TEXT)
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    for column in range(1, end_column + 1):
        worksheet.cell(row=3, column=column).fill = PatternFill("solid", fgColor=_LIGHT_GREEN)


def _append_headers(
    worksheet: Worksheet,
    *,
    row: int,
    headers: tuple[str, ...],
    fill_color: str = _DARK_GREEN,
    font_color: str = _WHITE,
) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.font = Font(name="Malgun Gothic", bold=True, color=font_color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _TABLE_BORDER
    worksheet.row_dimensions[row].height = 24


def _style_data_row(
    worksheet: Worksheet,
    *,
    row: int,
    columns: int,
    fill_color: str | None = None,
) -> None:
    for column in range(1, columns + 1):
        cell = worksheet.cell(row=row, column=column)
        cell.font = Font(name="Malgun Gothic", size=10, color=_DARK_TEXT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _TABLE_BORDER
        if fill_color is not None:
            cell.fill = PatternFill("solid", fgColor=fill_color)
    worksheet.row_dimensions[row].height = 23


def _finish_table(worksheet: Worksheet, *, header_row: int, column_count: int, table_name: str) -> None:
    end_row = max(header_row, worksheet.max_row)
    table_ref = f"A{header_row}:{_column_letter(column_count)}{end_row}"
    if end_row == header_row:
        worksheet.auto_filter.ref = table_ref
        return
    table = Table(displayName=table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _set_widths(worksheet: Worksheet, widths: tuple[float, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[_column_letter(index)].width = width


def _pad(values: tuple[str, ...], length: int) -> tuple[str, ...]:
    return values[:length] + ("",) * max(0, length - len(values))


def _format_datetime(value: datetime | None, *, date_only: bool = False) -> str:
    if value is None:
        return ""
    if value.tzinfo is None or value.utcoffset() is None:
        raise ValueError("workbook timestamp must include timezone information")
    aware = value.astimezone(KST)
    return aware.strftime("%Y-%m-%d" if date_only else "%Y-%m-%d %H:%M KST")


def _safe_cell(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
