from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from umacircle_bot.domain.errors import LegacyImportError
from umacircle_bot.domain.imports import normalize_import_sheet_name

DEFAULT_RATING_RULE_SHEET_NAME = "Rate 기준표"


@dataclass(frozen=True)
class WorkbookRatingRuleRow:
    grade: str
    participant_count: int
    converted_rank: int
    base_delta: Decimal


@dataclass(frozen=True)
class RatingRuleWorkbookPreparation:
    source_checksum: str
    sheet_name: str
    source_range: str
    rules: tuple[WorkbookRatingRuleRow, ...]


def prepare_rating_rule_workbook(
    path: str | Path,
    *,
    sheet_name: str = DEFAULT_RATING_RULE_SHEET_NAME,
) -> RatingRuleWorkbookPreparation:
    workbook_path = Path(path)
    normalized_sheet_name = normalize_import_sheet_name(sheet_name)
    source_checksum = _sha256_file(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        if normalized_sheet_name not in workbook.sheetnames:
            raise LegacyImportError(f"rating rule workbook is missing sheet: {normalized_sheet_name}")
        rules, source_range = _parse_rating_rule_sheet(workbook[normalized_sheet_name])
    finally:
        workbook.close()
    return RatingRuleWorkbookPreparation(
        source_checksum=source_checksum,
        sheet_name=normalized_sheet_name,
        source_range=source_range,
        rules=rules,
    )


def _parse_rating_rule_sheet(worksheet: Any) -> tuple[tuple[WorkbookRatingRuleRow, ...], str]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise LegacyImportError("rating rule worksheet is empty")

    parsed: list[WorkbookRatingRuleRow] = []
    index = 0
    while index < len(rows):
        row = rows[index]
        grade = _read_grade_header(row)
        if grade is None:
            index += 1
            continue
        if index + 2 >= len(rows):
            raise LegacyImportError(f"rating rule grade block {grade} is incomplete")
        participant_counts = _read_participant_counts(rows[index + 1], grade=grade)
        index += 2
        while index < len(rows):
            data_row = rows[index]
            if _read_grade_header(data_row) is not None:
                break
            converted_rank = _as_positive_int(_cell(data_row, 1))
            if converted_rank is not None:
                for column_index, participant_count in participant_counts.items():
                    value = _cell(data_row, column_index)
                    if value is not None:
                        parsed.append(
                            WorkbookRatingRuleRow(
                                grade=grade,
                                participant_count=participant_count,
                                converted_rank=converted_rank,
                                base_delta=_as_decimal(value),
                            )
                        )
            index += 1

    max_used_row = 0
    max_used_column = 0
    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            if value is not None:
                max_used_row = max(max_used_row, row_number)
                max_used_column = max(max_used_column, column_number)
    if not parsed:
        raise LegacyImportError("rating rule worksheet contains no numeric rules")
    keys = {(rule.grade, rule.participant_count, rule.converted_rank) for rule in parsed}
    if len(keys) != len(parsed):
        raise LegacyImportError("rating rule worksheet contains duplicate grade/participant/rank rules")
    return (
        tuple(sorted(parsed, key=lambda item: (item.grade, item.participant_count, item.converted_rank))),
        f"A1:{_column_name(max_used_column)}{max_used_row}",
    )


def _read_grade_header(row: tuple[Any, ...]) -> str | None:
    value = _cell(row, 0)
    if not isinstance(value, str) or _cell(row, 2) != "레이스 인원":
        return None
    grade = value.strip().upper()
    if not grade:
        return None
    return grade


def _read_participant_counts(row: tuple[Any, ...], *, grade: str) -> dict[int, int]:
    counts: dict[int, int] = {}
    for column_index in range(2, len(row)):
        participant_count = _as_positive_int(row[column_index])
        if participant_count is not None:
            counts[column_index] = participant_count
    if not counts:
        raise LegacyImportError(f"rating rule grade block {grade} has no participant counts")
    return counts


def _cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _as_positive_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, float) and value.is_integer() and value > 0:
        return int(value)
    return None


def _as_decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise LegacyImportError("rating rule base delta must be numeric")
    return Decimal(str(value))


def _column_name(column_number: int) -> str:
    if column_number <= 0:
        raise LegacyImportError("rating rule worksheet has no populated cells")
    letters = ""
    remaining = column_number
    while remaining:
        remaining, offset = divmod(remaining - 1, 26)
        letters = chr(ord("A") + offset) + letters
    return letters


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
