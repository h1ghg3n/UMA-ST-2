from __future__ import annotations

from dataclasses import asdict, dataclass
from decimal import ROUND_HALF_UP, Decimal
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from umacircle_bot.sheets.legacy_room_ledger import (
    LegacyRoomLedgerReconciliation,
    parse_legacy_room_ledger_rows,
    parse_legacy_room_payout_rows,
    parse_legacy_room_race_rows,
    reconcile_legacy_room_ledger,
)
from umacircle_bot.sheets.row_parsers import RowParseIssue, parse_room_point_rows
from umacircle_bot.sheets.xlsx_dry_run import WorkbookKind, inspect_xlsx_workbook

RECONCILIATION_SHEETS = ("베팅 Data", "룸매치 Data", "정답배당", "동친 포인트")


@dataclass(frozen=True)
class LegacyRoomDryRunReport:
    source_name: str
    source_checksum: str
    ledger_row_count: int
    point_row_count: int
    race_row_count: int
    payout_row_count: int
    movement_count: int
    final_point_total: int | None
    payout_correction_total: int
    corrected_final_point_total: int | None
    issues: tuple[dict[str, Any], ...]
    payout_mismatches: tuple[dict[str, Any], ...]
    payout_correction_plan: tuple[dict[str, Any], ...]

    @property
    def has_errors(self) -> bool:
        return any(issue["severity"] == "error" for issue in self.issues)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def build_legacy_room_dry_run_report(path: str | Path) -> LegacyRoomDryRunReport:
    workbook_path = Path(path).expanduser().resolve()
    structure = inspect_xlsx_workbook(workbook_path, kind=WorkbookKind.ROOM_MATCH, sheets=RECONCILIATION_SHEETS)
    issues: list[dict[str, Any]] = [asdict(issue) for issue in structure.issues]
    missing = set(RECONCILIATION_SHEETS) - set(structure.sheet_names)
    if missing:
        return _empty_report(workbook_path, issues)

    workbook = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    try:
        ledger_report = parse_legacy_room_ledger_rows(_rows(workbook["베팅 Data"]))
        race_report = parse_legacy_room_race_rows(_rows(workbook["룸매치 Data"]))
        payout_report = parse_legacy_room_payout_rows(_rows(workbook["정답배당"]))
        point_report = parse_room_point_rows(_rows(workbook["동친 포인트"]))
    finally:
        workbook.close()

    for report in (ledger_report, race_report, payout_report, point_report):
        issues.extend(_parse_issue(issue) for issue in report.issues)

    reconciliation: LegacyRoomLedgerReconciliation | None = None
    if not any(issue["severity"] == "error" for issue in issues):
        reconciliation = reconcile_legacy_room_ledger(
            ledger_report.rows,
            point_report.rows,
            race_report.rows,
            payout_report.rows,
        )
        issues.extend(asdict(issue) for issue in reconciliation.issues)

    mismatch_rows = {
        issue["row_number"]
        for issue in issues
        if issue["code"] == "payout_rate_mismatch" and issue.get("row_number") is not None
    }
    ledger_by_row = {row.row_number: row for row in ledger_report.rows}
    participant_counts = {row.race_label: row.participant_count for row in race_report.rows}
    payouts_by_key = {(row.race_label, row.bet_type, row.numbers): row.payout_rate for row in payout_report.rows}
    payout_mismatches = tuple(
        {
            "row_number": row.row_number,
            "race_label": row.race_label,
            "participant_name": row.participant_name,
            "bet_type": row.bet_type,
            "numbers": list(row.numbers),
            "amount": row.amount,
            "cached_payout_rate": str(row.payout_rate),
            "cached_payout_amount": row.payout_amount,
            "expected_payout_rate": str(expected_rate),
            "expected_payout_amount": (expected_amount := _payout_amount(row.amount, expected_rate)),
            "payout_shortfall": expected_amount - row.payout_amount,
            "participant_count": participant_counts.get(row.race_label),
        }
        for row_number in sorted(mismatch_rows)
        if (row := ledger_by_row.get(row_number)) is not None
        and (expected_rate := payouts_by_key.get((row.race_label, row.bet_type, row.numbers))) is not None
    )
    correction_by_participant: dict[str, int] = {}
    for mismatch in payout_mismatches:
        correction_by_participant[mismatch["participant_name"]] = (
            correction_by_participant.get(mismatch["participant_name"], 0) + mismatch["payout_shortfall"]
        )
    points_by_name = {row.nickname: row for row in point_report.rows}
    correction_plan = tuple(
        {
            "participant_name": name,
            "discord_user_id": points_by_name[name].discord_user_id,
            "current_balance": points_by_name[name].points,
            "correction_amount": amount,
            "corrected_balance": points_by_name[name].points + amount,
        }
        for name, amount in sorted(correction_by_participant.items())
        if name in points_by_name
    )
    correction_total = sum(correction_by_participant.values())
    final_point_total = (
        sum(identity.current_balance for identity in reconciliation.identities) if reconciliation is not None else None
    )
    return LegacyRoomDryRunReport(
        source_name=workbook_path.name,
        source_checksum=_file_checksum(workbook_path),
        ledger_row_count=len(ledger_report.rows),
        point_row_count=len(point_report.rows),
        race_row_count=len(race_report.rows),
        payout_row_count=len(payout_report.rows),
        movement_count=len(reconciliation.movements) if reconciliation is not None else 0,
        final_point_total=final_point_total,
        payout_correction_total=correction_total,
        corrected_final_point_total=(final_point_total + correction_total if final_point_total is not None else None),
        issues=tuple(issues),
        payout_mismatches=payout_mismatches,
        payout_correction_plan=correction_plan,
    )


def _empty_report(path: Path, issues: list[dict[str, Any]]) -> LegacyRoomDryRunReport:
    return LegacyRoomDryRunReport(
        source_name=path.name,
        source_checksum=_file_checksum(path),
        ledger_row_count=0,
        point_row_count=0,
        race_row_count=0,
        payout_row_count=0,
        movement_count=0,
        final_point_total=None,
        payout_correction_total=0,
        corrected_final_point_total=None,
        issues=tuple(issues),
        payout_mismatches=(),
        payout_correction_plan=(),
    )


def _rows(worksheet: Any):
    return worksheet.iter_rows(values_only=True)


def _parse_issue(issue: RowParseIssue) -> dict[str, Any]:
    return {
        "severity": "error",
        "code": issue.code,
        "message": issue.message,
        "row_number": issue.row_number,
        "sheet_name": issue.sheet_name,
    }


def _file_checksum(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _payout_amount(stake: int, payout_rate: Decimal) -> int:
    return int((Decimal(stake) * payout_rate).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
