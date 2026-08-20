from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from umacircle_bot.domain.errors import LegacyImportError
from umacircle_bot.services.legacy_import import (
    LegacyIdentityPointImportPlan,
    build_legacy_identity_point_import_plan,
)
from umacircle_bot.sheets.legacy_result_plan import (
    LegacyRoomResultImportPlan,
    build_legacy_room_result_import_plan,
)
from umacircle_bot.sheets.legacy_room_ledger import (
    LegacyRoomLedgerReconciliation,
    LegacyRoomLedgerRow,
    LegacyRoomPayoutSourceRow,
    LegacyRoomRaceSourceRow,
    parse_legacy_room_ledger_rows,
    parse_legacy_room_payout_rows,
    parse_legacy_room_race_rows,
    reconcile_legacy_room_ledger,
)
from umacircle_bot.sheets.row_parsers import (
    MatchDataSourceRow,
    RoomPointSourceRow,
    parse_match_data_rows,
    parse_room_point_rows,
)
from umacircle_bot.sheets.xlsx_dry_run import normalize_xlsx_workbook_path

MAX_LEGACY_WORKBOOK_SIZE_BYTES = 512 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 10_000
MAX_XLSX_EXPANDED_SIZE_BYTES = 2 * 1024 * 1024 * 1024
REQUIRED_IDENTITY_POINT_SHEETS = (
    "베팅 Data",
    "동친 포인트",
    "룸매치 Data",
    "정답배당",
)
REQUIRED_RESULT_SHEETS = (
    "Data",
    "룸매치 Data",
    "정답배당",
)


@dataclass(frozen=True)
class LegacyIdentityPointWorkbookPreparation:
    workbook_path: Path
    source_checksum: str
    plan: LegacyIdentityPointImportPlan
    ledger_row_count: int
    race_row_count: int
    payout_row_count: int
    warning_codes: tuple[str, ...]
    race_rows: tuple[LegacyRoomRaceSourceRow, ...]
    ledger_rows: tuple[LegacyRoomLedgerRow, ...]
    point_rows: tuple[RoomPointSourceRow, ...]
    reconciliation: LegacyRoomLedgerReconciliation


@dataclass(frozen=True)
class LegacyRoomResultWorkbookPreparation:
    workbook_path: Path
    source_checksum: str
    plan: LegacyRoomResultImportPlan
    data_row_count: int
    race_row_count: int
    payout_row_count: int
    data_rows: tuple[MatchDataSourceRow, ...]
    race_rows: tuple[LegacyRoomRaceSourceRow, ...]
    payout_rows: tuple[LegacyRoomPayoutSourceRow, ...]


def prepare_legacy_identity_point_workbook(
    path: str | Path,
    *,
    source_identifier: str,
    max_file_size_bytes: int = MAX_LEGACY_WORKBOOK_SIZE_BYTES,
) -> LegacyIdentityPointWorkbookPreparation:
    workbook_path = normalize_xlsx_workbook_path(path)
    _validate_max_file_size(max_file_size_bytes)
    if workbook_path.stat().st_size > max_file_size_bytes:
        raise LegacyImportError("legacy workbook exceeds the configured size limit")
    _validate_xlsx_archive(workbook_path)
    checksum_before = calculate_file_sha256(workbook_path)

    try:
        workbook = load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise LegacyImportError("legacy workbook could not be opened safely") from exc
    try:
        missing_sheets = sorted(set(REQUIRED_IDENTITY_POINT_SHEETS) - set(workbook.sheetnames))
        if missing_sheets:
            raise LegacyImportError(f"legacy workbook is missing required sheet: {missing_sheets[0]}")
        ledger_report = parse_legacy_room_ledger_rows(workbook["베팅 Data"].iter_rows(values_only=True))
        point_report = parse_room_point_rows(workbook["동친 포인트"].iter_rows(values_only=True))
        race_report = parse_legacy_room_race_rows(workbook["룸매치 Data"].iter_rows(values_only=True))
        payout_report = parse_legacy_room_payout_rows(workbook["정답배당"].iter_rows(values_only=True))
    finally:
        workbook.close()

    parse_issue_counts = {
        "ledger": len(ledger_report.issues),
        "points": len(point_report.issues),
        "races": len(race_report.issues),
        "payouts": len(payout_report.issues),
    }
    if issue_count := sum(parse_issue_counts.values()):
        affected = min(name for name, count in parse_issue_counts.items() if count)
        raise LegacyImportError(f"legacy workbook has {issue_count} parse issue(s); first affected area: {affected}")

    reconciliation = reconcile_legacy_room_ledger(
        ledger_report.rows,
        point_report.rows,
        race_report.rows,
        payout_report.rows,
    )
    if reconciliation.has_errors:
        first_error = next(issue for issue in reconciliation.issues if issue.severity == "error")
        raise LegacyImportError(f"legacy workbook reconciliation failed: {first_error.code}")
    plan = build_legacy_identity_point_import_plan(
        point_report.rows,
        reconciliation.identities,
        source_identifier=source_identifier,
    )

    checksum_after = calculate_file_sha256(workbook_path)
    if checksum_before != checksum_after:
        raise LegacyImportError("legacy workbook changed while it was being inspected")
    return LegacyIdentityPointWorkbookPreparation(
        workbook_path=workbook_path,
        source_checksum=checksum_after,
        plan=plan,
        ledger_row_count=len(ledger_report.rows),
        race_row_count=len(race_report.rows),
        payout_row_count=len(payout_report.rows),
        warning_codes=tuple(issue.code for issue in reconciliation.issues if issue.severity == "warning"),
        race_rows=race_report.rows,
        ledger_rows=ledger_report.rows,
        point_rows=point_report.rows,
        reconciliation=reconciliation,
    )


def prepare_legacy_room_result_workbook(
    path: str | Path,
    *,
    source_identifier: str,
    source_utc_offset_minutes: int,
    max_file_size_bytes: int = MAX_LEGACY_WORKBOOK_SIZE_BYTES,
) -> LegacyRoomResultWorkbookPreparation:
    workbook_path = normalize_xlsx_workbook_path(path)
    _validate_max_file_size(max_file_size_bytes)
    if workbook_path.stat().st_size > max_file_size_bytes:
        raise LegacyImportError("legacy workbook exceeds the configured size limit")
    _validate_xlsx_archive(workbook_path)
    checksum_before = calculate_file_sha256(workbook_path)

    workbook = None
    try:
        workbook = load_workbook(
            filename=workbook_path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        missing_sheets = sorted(set(REQUIRED_RESULT_SHEETS) - set(workbook.sheetnames))
        if missing_sheets:
            raise LegacyImportError(f"legacy workbook is missing required sheet: {missing_sheets[0]}")
        data_report = parse_match_data_rows(workbook["Data"].iter_rows(values_only=True))
        race_report = parse_legacy_room_race_rows(workbook["룸매치 Data"].iter_rows(values_only=True))
        payout_report = parse_legacy_room_payout_rows(workbook["정답배당"].iter_rows(values_only=True))
    except LegacyImportError:
        raise
    except Exception as exc:
        raise LegacyImportError("legacy workbook could not be opened safely") from exc
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception as exc:
                raise LegacyImportError("legacy workbook could not be opened safely") from exc

    parse_issue_counts = {
        "data": len(data_report.issues),
        "races": len(race_report.issues),
        "payouts": len(payout_report.issues),
    }
    if issue_count := sum(parse_issue_counts.values()):
        affected = min(name for name, count in parse_issue_counts.items() if count)
        raise LegacyImportError(f"legacy workbook has {issue_count} parse issue(s); first affected area: {affected}")
    plan = build_legacy_room_result_import_plan(
        data_report.rows,
        race_report.rows,
        payout_report.rows,
        source_identifier=source_identifier,
        source_utc_offset_minutes=source_utc_offset_minutes,
    )

    checksum_after = calculate_file_sha256(workbook_path)
    if checksum_before != checksum_after:
        raise LegacyImportError("legacy workbook changed while it was being inspected")
    return LegacyRoomResultWorkbookPreparation(
        workbook_path=workbook_path,
        source_checksum=checksum_after,
        plan=plan,
        data_row_count=len(data_report.rows),
        race_row_count=len(race_report.rows),
        payout_row_count=len(payout_report.rows),
        data_rows=data_report.rows,
        race_rows=race_report.rows,
        payout_rows=payout_report.rows,
    )


def calculate_file_sha256(path: str | Path, *, chunk_size: int = 1024 * 1024) -> str:
    file_path = Path(path)
    if not isinstance(chunk_size, int) or isinstance(chunk_size, bool) or not 4096 <= chunk_size <= 16 * 1024 * 1024:
        raise ValueError("chunk_size must be an integer between 4096 and 16777216")
    digest = sha256()
    with file_path.open("rb") as source_file:
        while chunk := source_file.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_max_file_size(value: int) -> None:
    if not isinstance(value, int) or isinstance(value, bool) or not 1 <= value <= 2 * 1024 * 1024 * 1024:
        raise ValueError("max_file_size_bytes must be an integer between 1 and 2147483648")


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
    except (BadZipFile, OSError) as exc:
        raise LegacyImportError("legacy workbook is not a valid XLSX archive") from exc
    if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
        raise LegacyImportError("legacy workbook archive contains too many entries")
    expanded_size = 0
    for entry in entries:
        member_path = PurePosixPath(entry.filename)
        if entry.flag_bits & 0x1:
            raise LegacyImportError("encrypted XLSX archive entries are not supported")
        if member_path.is_absolute() or ".." in member_path.parts:
            raise LegacyImportError("legacy workbook archive contains an unsafe member path")
        expanded_size += entry.file_size
        if expanded_size > MAX_XLSX_EXPANDED_SIZE_BYTES:
            raise LegacyImportError("legacy workbook expanded data exceeds the safety limit")
