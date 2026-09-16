"""MariaDB evidence for the read-only current Circle Point workbook export."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.engine import Engine

from uma_st2.compose import compose_circle_point_exports
from uma_st2.infrastructure.database import DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    CirclePointORM,
    OperationORM,
    PersonaORM,
    PointTransactionORM,
)

pytestmark = pytest.mark.integration


@dataclass(frozen=True, slots=True)
class SeededExport:
    persona_id: str
    operation_id: int
    transaction_ids: tuple[int, int]


def _seed(engine: Engine) -> SeededExport:
    now = datetime.now(UTC).replace(tzinfo=None)
    persona_id = str(uuid4())
    with engine.begin() as connection:
        connection.execute(
            PersonaORM.__table__.insert().values(
                id=persona_id,
                display_name="MariaDB Circle Point Export",
                status="warning",
                created_at=now,
                updated_at=now,
            )
        )
        connection.execute(
            CirclePointORM.__table__.insert().values(
                persona_id=persona_id,
                balance=777,
                updated_at=now,
            )
        )
        operation_id = connection.execute(
            OperationORM.__table__.insert().values(
                guild_id=None,
                correlation_id=None,
                actor_discord_user_id=None,
                idempotency_key=None,
                request_fingerprint=None,
                reason="retained MariaDB reason",
                created_at=now,
            )
        ).inserted_primary_key[0]
        first_id = connection.execute(
            PointTransactionORM.__table__.insert().values(
                persona_id=persona_id,
                operation_id=operation_id,
                action="manual_grant",
                amount=500,
                created_at=now,
            )
        ).inserted_primary_key[0]
        second_id = connection.execute(
            PointTransactionORM.__table__.insert().values(
                persona_id=persona_id,
                operation_id=None,
                action="match_bet_stake",
                amount=-10,
                created_at=now,
            )
        ).inserted_primary_key[0]
    return SeededExport(persona_id, operation_id, (first_id, second_id))


def _cleanup(engine: Engine, seeded: SeededExport) -> None:
    with engine.begin() as connection:
        connection.execute(delete(PointTransactionORM).where(PointTransactionORM.id.in_(seeded.transaction_ids)))
        connection.execute(delete(CirclePointORM).where(CirclePointORM.persona_id == seeded.persona_id))
        connection.execute(delete(OperationORM).where(OperationORM.id == seeded.operation_id))
        connection.execute(delete(PersonaORM).where(PersonaORM.id == seeded.persona_id))


def test_mariadb_circle_point_export_preserves_current_balance_and_retained_rows_without_writes(
    migrated_engine: Engine,
) -> None:
    seeded = _seed(migrated_engine)
    runtime = DatabaseRuntime.from_engine(migrated_engine)
    try:
        with migrated_engine.connect() as connection:
            before = connection.scalar(
                select(func.count())
                .select_from(PointTransactionORM)
                .where(PointTransactionORM.id.in_(seeded.transaction_ids))
            )

        artifact = compose_circle_point_exports(runtime).export_current()

        with migrated_engine.connect() as connection:
            after = connection.scalar(
                select(func.count())
                .select_from(PointTransactionORM)
                .where(PointTransactionORM.id.in_(seeded.transaction_ids))
            )
        assert before == after == 2
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            wallet_rows = list(workbook["현재 잔액"].iter_rows(min_row=2, values_only=True))
            transaction_rows = list(workbook["거래 내역"].iter_rows(min_row=2, values_only=True))
            wallet = next(row for row in wallet_rows if row[0] == seeded.persona_id)
            transactions = [row for row in transaction_rows if row[0] in seeded.transaction_ids]
            assert wallet[1:4] == ("MariaDB Circle Point Export", "warning", 777)
            assert [row[0] for row in transactions] == list(seeded.transaction_ids)
            assert transactions[0][6:] == (seeded.operation_id, "retained MariaDB reason")
            assert transactions[1][6:] == (None, None)
        finally:
            workbook.close()
    finally:
        runtime.dispose()
        _cleanup(migrated_engine, seeded)
