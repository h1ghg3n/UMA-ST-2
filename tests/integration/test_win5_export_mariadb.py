"""MariaDB evidence for complete read-only WIN5 Season workbook export."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.engine import Engine

from uma_st2.compose import compose_win5_season_exports
from uma_st2.infrastructure.database import DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    PersonaORM,
    Win5RaceORM,
    Win5ResultORM,
    Win5RoundORM,
    Win5ScoreEventItemORM,
    Win5ScoreEventORM,
    Win5ScoreORM,
    Win5SeasonORM,
    Win5SubmissionORM,
    Win5SubmissionPickORM,
)

pytestmark = pytest.mark.integration


@dataclass(frozen=True, slots=True)
class SeededExport:
    season_id: int
    round_id: int
    race_ids: tuple[int, int]
    result_id: int
    persona_id: str
    submission_id: int
    pick_ids: tuple[int, int]
    event_id: int


def _seed(engine: Engine) -> SeededExport:
    now = datetime.now(UTC).replace(tzinfo=None)
    suffix = uuid4().hex
    persona_id = str(uuid4())
    with engine.begin() as connection:
        connection.execute(
            PersonaORM.__table__.insert().values(
                id=persona_id,
                display_name=f"Export {suffix}",
                status="normal",
                created_at=now,
                updated_at=now,
            )
        )
        season_id = connection.execute(
            Win5SeasonORM.__table__.insert().values(
                name=f"Closed Export {suffix}",
                status="closed",
                active_marker=None,
                starts_at=now,
                ends_at=now,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        round_id = connection.execute(
            Win5RoundORM.__table__.insert().values(
                season_id=season_id,
                type="special",
                status="scored",
                name=f"Mixed Void {suffix}",
                opens_at=now,
                closes_at=now,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        void_race_id = connection.execute(
            Win5RaceORM.__table__.insert().values(
                round_id=round_id,
                name="Void Race",
                scheduled_at=now,
                void_reason="official no-contest",
                voided_at=now,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        live_race_id = connection.execute(
            Win5RaceORM.__table__.insert().values(
                round_id=round_id,
                name="Live Race",
                scheduled_at=now,
                void_reason=None,
                voided_at=None,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        result_id = connection.execute(
            Win5ResultORM.__table__.insert().values(
                race_id=live_race_id,
                race_entry_id=None,
                gate_number=7,
                position=1,
                created_at=now,
            )
        ).inserted_primary_key[0]
        submission_id = connection.execute(
            Win5SubmissionORM.__table__.insert().values(
                round_id=round_id,
                persona_id=persona_id,
                tier="SPECIAL_WINNER",
                status="accepted",
                active_marker=True,
                version=1,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        pick_ids = tuple(
            connection.execute(
                Win5SubmissionPickORM.__table__.insert().values(
                    submission_id=submission_id,
                    race_id=race_id,
                    race_entry_id=None,
                    gate_number=gate_number,
                    position=1,
                )
            ).inserted_primary_key[0]
            for race_id, gate_number in ((void_race_id, 3), (live_race_id, 7))
        )
        event_id = connection.execute(
            Win5ScoreEventORM.__table__.insert().values(
                operation_id=None,
                season_id=season_id,
                round_id=round_id,
                race_id=None,
                submission_id=submission_id,
                submission_version=1,
                persona_id=persona_id,
                tier="SPECIAL_WINNER",
                result_fingerprint="b" * 64,
                scoring_policy_version="special/v1",
                reward_policy_version="reward/v1",
                exact_count=1,
                wrong_position_count=0,
                off_board_count=0,
                missing_count=0,
                season_score_delta=1,
                top1_score_delta=1,
                circle_point_reward=0,
                created_at=now,
            )
        ).inserted_primary_key[0]
        connection.execute(
            Win5ScoreEventItemORM.__table__.insert(),
            [
                {
                    "score_event_id": event_id,
                    "race_id": void_race_id,
                    "position": 1,
                    "submission_pick_id": pick_ids[0],
                    "matched_result_id": None,
                    "outcome": "void",
                    "season_score_delta": 0,
                },
                {
                    "score_event_id": event_id,
                    "race_id": live_race_id,
                    "position": 1,
                    "submission_pick_id": pick_ids[1],
                    "matched_result_id": result_id,
                    "outcome": "exact",
                    "season_score_delta": 1,
                },
            ],
        )
        connection.execute(
            Win5ScoreORM.__table__.insert().values(
                season_id=season_id,
                persona_id=persona_id,
                season_score=1,
                top1_score=1,
                updated_at=now,
            )
        )
    return SeededExport(
        season_id=season_id,
        round_id=round_id,
        race_ids=(void_race_id, live_race_id),
        result_id=result_id,
        persona_id=persona_id,
        submission_id=submission_id,
        pick_ids=pick_ids,
        event_id=event_id,
    )


def _cleanup(engine: Engine, seeded: SeededExport) -> None:
    with engine.begin() as connection:
        connection.execute(delete(Win5ScoreEventItemORM).where(Win5ScoreEventItemORM.score_event_id == seeded.event_id))
        connection.execute(delete(Win5ScoreEventORM).where(Win5ScoreEventORM.id == seeded.event_id))
        connection.execute(delete(Win5ScoreORM).where(Win5ScoreORM.season_id == seeded.season_id))
        connection.execute(delete(Win5SubmissionPickORM).where(Win5SubmissionPickORM.id.in_(seeded.pick_ids)))
        connection.execute(delete(Win5SubmissionORM).where(Win5SubmissionORM.id == seeded.submission_id))
        connection.execute(delete(Win5ResultORM).where(Win5ResultORM.id == seeded.result_id))
        connection.execute(delete(Win5RaceORM).where(Win5RaceORM.id.in_(seeded.race_ids)))
        connection.execute(delete(Win5RoundORM).where(Win5RoundORM.id == seeded.round_id))
        connection.execute(delete(Win5SeasonORM).where(Win5SeasonORM.id == seeded.season_id))
        connection.execute(delete(PersonaORM).where(PersonaORM.id == seeded.persona_id))


def test_mariadb_export_reads_mixed_void_snapshot_without_canonical_writes(
    migrated_engine: Engine,
) -> None:
    seeded = _seed(migrated_engine)
    runtime = DatabaseRuntime.from_engine(migrated_engine)
    try:
        with migrated_engine.connect() as connection:
            before = connection.scalar(
                select(func.count()).select_from(Win5ScoreEventORM).where(Win5ScoreEventORM.id == seeded.event_id)
            )

        exports = compose_win5_season_exports(runtime)
        assert seeded.season_id in {choice.id for choice in exports.search_seasons(query="Closed Export")}
        artifact = exports.export_season(season_id=seeded.season_id)

        with migrated_engine.connect() as connection:
            after = connection.scalar(
                select(func.count()).select_from(Win5ScoreEventORM).where(Win5ScoreEventORM.id == seeded.event_id)
            )
        assert before == after == 1
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            rows = list(workbook["제출 및 판정"].iter_rows(min_row=2, values_only=True))
            assert [(row[12], row[14], row[15], row[16]) for row in rows] == [
                ("Void Race", "3", "void", "void"),
                ("Live Race", "7", "7", "exact"),
            ]
        finally:
            workbook.close()
    finally:
        _cleanup(migrated_engine, seeded)
