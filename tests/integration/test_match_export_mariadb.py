"""MariaDB evidence for complete read-only Circle Match Season workbook export."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.engine import Engine

from uma_st2.compose import compose_match_season_exports
from uma_st2.infrastructure.database import DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    GameAccountORM,
    MatchEntryORM,
    MatchORM,
    PersonaORM,
    StadiumCourseORM,
    StadiumORM,
    UmamusumeORM,
)

pytestmark = pytest.mark.integration


@dataclass(frozen=True, slots=True)
class SeededExport:
    persona_id: str
    game_account_id: int
    character_id: int
    stadium_id: int
    course_id: int
    match_id: int
    entry_id: int


def _seed(engine: Engine) -> SeededExport:
    now = datetime.now(UTC).replace(tzinfo=None)
    suffix = uuid4().hex
    external_seed = int(suffix[:8], 16)
    persona_id = str(uuid4())
    with engine.begin() as connection:
        connection.execute(
            PersonaORM.__table__.insert().values(
                id=persona_id,
                display_name=f"Match Export {suffix}",
                status="normal",
                created_at=now,
                updated_at=now,
            )
        )
        game_account_id = connection.execute(
            GameAccountORM.__table__.insert().values(
                persona_id=persona_id,
                game_region="KR",
                uma_pid=None,
                nickname=f"Account {suffix}",
                affiliation="Circle",
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        character_id = connection.execute(
            UmamusumeORM.__table__.insert().values(
                external_id=1_000_000_000 + external_seed,
                name_jp=f"Horse {suffix}",
                name_ko=None,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        stadium_id = connection.execute(
            StadiumORM.__table__.insert().values(
                external_id=2_000_000_000 + external_seed,
                name_jp=f"Stadium {suffix}",
                name_ko=None,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        course_id = connection.execute(
            StadiumCourseORM.__table__.insert().values(
                stadium_id=stadium_id,
                external_id=3_000_000_000 + external_seed,
                surface="turf",
                distance=1600,
                direction="right",
                layout="standard",
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        match_id = connection.execute(
            MatchORM.__table__.insert().values(
                name=f"Imported Export {suffix}",
                description=None,
                source_kind="imported_v1",
                grade="LISTED",
                stadium_course_id=course_id,
                scheduled_at=datetime(2098, 7, 1),
                status="result_confirmed",
                terminal_reason=None,
                finish_time_ms=100000,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
        entry_id = connection.execute(
            MatchEntryORM.__table__.insert().values(
                match_id=match_id,
                game_account_id=game_account_id,
                owner_at_event_persona_id=persona_id,
                affiliation_at_event="Historical Circle",
                umamusume_id=character_id,
                umamusume_variant_id=None,
                entry_number=7,
                running_style=None,
                training_grade=None,
                rank=1,
                popularity_rank=None,
                margin=None,
                created_at=now,
                updated_at=now,
            )
        ).inserted_primary_key[0]
    return SeededExport(
        persona_id=persona_id,
        game_account_id=game_account_id,
        character_id=character_id,
        stadium_id=stadium_id,
        course_id=course_id,
        match_id=match_id,
        entry_id=entry_id,
    )


def _cleanup(engine: Engine, seeded: SeededExport) -> None:
    with engine.begin() as connection:
        connection.execute(delete(MatchEntryORM).where(MatchEntryORM.id == seeded.entry_id))
        connection.execute(delete(MatchORM).where(MatchORM.id == seeded.match_id))
        connection.execute(delete(StadiumCourseORM).where(StadiumCourseORM.id == seeded.course_id))
        connection.execute(delete(StadiumORM).where(StadiumORM.id == seeded.stadium_id))
        connection.execute(delete(UmamusumeORM).where(UmamusumeORM.id == seeded.character_id))
        connection.execute(delete(GameAccountORM).where(GameAccountORM.id == seeded.game_account_id))
        connection.execute(delete(PersonaORM).where(PersonaORM.id == seeded.persona_id))


def test_mariadb_match_export_reads_imported_raw_snapshot_without_canonical_writes(
    migrated_engine: Engine,
) -> None:
    seeded = _seed(migrated_engine)
    runtime = DatabaseRuntime.from_engine(migrated_engine)
    try:
        with migrated_engine.connect() as connection:
            before = connection.scalar(
                select(func.count()).select_from(MatchEntryORM).where(MatchEntryORM.id == seeded.entry_id)
            )

        exports = compose_match_season_exports(runtime)
        choices = exports.search_seasons(query="2098")
        assert [(choice.key, choice.name) for choice in choices] == [("2098-split-2", "2098 Split 2")]
        artifact = exports.export_season(season_key="2098-split-2")

        with migrated_engine.connect() as connection:
            after = connection.scalar(
                select(func.count()).select_from(MatchEntryORM).where(MatchEntryORM.id == seeded.entry_id)
            )
        assert before == after == 1
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            match_rows = list(workbook["경기"].iter_rows(min_row=2, values_only=True))
            entry_rows = list(workbook["출전 및 결과"].iter_rows(min_row=2, values_only=True))
            assert len(match_rows) == len(entry_rows) == 1
            assert match_rows[0][3] == "imported_v1"
            assert entry_rows[0][6] == 7
            assert workbook["베팅 및 정산"].max_row == 1
            assert workbook["레이팅"].max_row == 1
        finally:
            workbook.close()
    finally:
        runtime.dispose()
        _cleanup(migrated_engine, seeded)
