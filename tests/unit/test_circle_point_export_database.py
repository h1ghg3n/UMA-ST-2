"""SQLAlchemy Circle Point export projection tests on a disposable local schema."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select

from uma_st2.application.exporting import CirclePointExportInvalidSourceError
from uma_st2.compose import compose_circle_point_exports
from uma_st2.infrastructure.database import Base, DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    CirclePointORM,
    OperationORM,
    PersonaORM,
    PointTransactionORM,
)

NOW = datetime(2026, 8, 29)


def _runtime() -> DatabaseRuntime:
    runtime = DatabaseRuntime.from_engine(create_engine("sqlite+pysqlite:///:memory:"))
    Base.metadata.create_all(runtime.engine)
    return runtime


def test_composed_export_reads_authoritative_wallets_and_nullable_operation_provenance() -> None:
    runtime = _runtime()
    try:
        with runtime.engine.begin() as connection:
            connection.execute(
                PersonaORM.__table__.insert(),
                [
                    {
                        "id": "persona-a",
                        "display_name": "Owner A",
                        "status": "normal",
                        "created_at": NOW,
                        "updated_at": NOW,
                    },
                    {
                        "id": "persona-b",
                        "display_name": "Owner B",
                        "status": "warning",
                        "created_at": NOW,
                        "updated_at": NOW,
                    },
                ],
            )
            connection.execute(
                CirclePointORM.__table__.insert(),
                [
                    {"persona_id": "persona-b", "balance": 70, "updated_at": NOW},
                    {"persona_id": "persona-a", "balance": 0, "updated_at": NOW},
                ],
            )
            connection.execute(
                OperationORM.__table__.insert().values(
                    id=20,
                    guild_id=None,
                    correlation_id=None,
                    actor_discord_user_id=None,
                    idempotency_key=None,
                    request_fingerprint=None,
                    reason="retained reason",
                    created_at=NOW,
                )
            )
            connection.execute(
                PointTransactionORM.__table__.insert(),
                [
                    {
                        "id": 11,
                        "persona_id": "persona-b",
                        "operation_id": None,
                        "action": "match_bet_stake",
                        "amount": -10,
                        "created_at": datetime(2026, 8, 29, 1),
                    },
                    {
                        "id": 10,
                        "persona_id": "persona-a",
                        "operation_id": 20,
                        "action": "manual_grant",
                        "amount": 50,
                        "created_at": datetime(2026, 8, 29),
                    },
                ],
            )
        with runtime.engine.connect() as connection:
            before = connection.scalar(select(func.count()).select_from(PointTransactionORM))

        artifact = compose_circle_point_exports(runtime).export_current()

        with runtime.engine.connect() as connection:
            after = connection.scalar(select(func.count()).select_from(PointTransactionORM))
        assert before == after == 2
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            wallet_rows = list(workbook["현재 잔액"].iter_rows(min_row=2, values_only=True))
            transaction_rows = list(workbook["거래 내역"].iter_rows(min_row=2, values_only=True))
            assert [row[0] for row in wallet_rows] == ["persona-a", "persona-b"]
            assert [row[0] for row in transaction_rows] == [10, 11]
            assert transaction_rows[0][6:] == (20, "retained reason")
            assert transaction_rows[1][6:] == (None, None)
        finally:
            workbook.close()
    finally:
        runtime.dispose()


def test_composed_export_fails_closed_when_retained_owner_has_no_wallet() -> None:
    runtime = _runtime()
    try:
        with runtime.engine.begin() as connection:
            connection.execute(
                PersonaORM.__table__.insert().values(
                    id="persona-missing-wallet",
                    display_name="Missing Wallet",
                    status="normal",
                    created_at=NOW,
                    updated_at=NOW,
                )
            )
            connection.execute(
                PointTransactionORM.__table__.insert().values(
                    id=10,
                    persona_id="persona-missing-wallet",
                    operation_id=None,
                    action="manual_grant",
                    amount=10,
                    created_at=NOW,
                )
            )

        with pytest.raises(CirclePointExportInvalidSourceError):
            compose_circle_point_exports(runtime).export_current()
    finally:
        runtime.dispose()
