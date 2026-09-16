"""SQLAlchemy Circle Match export projection tests on a disposable local schema."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select

from uma_st2.application.match import (
    CancelledMatch,
    MatchCancellationRefund,
    MatchSettlementAppliedOdds,
    MatchSettlementPayout,
    MatchSettlementRating,
    MatchSettlementResultAuthority,
    SettledMatch,
    StoredMatchRefundPublication,
)
from uma_st2.compose import compose_match_season_exports
from uma_st2.domain.betting import BetType
from uma_st2.domain.match import MatchStatus
from uma_st2.domain.publication import PublicationStatus
from uma_st2.infrastructure.database import Base, DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    BetORM,
    GameAccountORM,
    MatchEntryORM,
    MatchOperationORM,
    MatchORM,
    OperationORM,
    PersonaORM,
    StadiumCourseORM,
    StadiumORM,
    UmamusumeORM,
)

NOW = datetime(2026, 8, 29)


def _seed(runtime: DatabaseRuntime) -> None:
    with runtime.engine.begin() as connection:
        connection.execute(
            PersonaORM.__table__.insert().values(
                id="persona-a",
                display_name="Owner",
                status="normal",
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            GameAccountORM.__table__.insert().values(
                id=1,
                persona_id="persona-a",
                game_region="KR",
                uma_pid=None,
                nickname="Account",
                affiliation="Circle",
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            UmamusumeORM.__table__.insert().values(
                id=1,
                external_id=101,
                name_jp="Horse JP",
                name_ko="Horse KO",
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            StadiumORM.__table__.insert().values(
                id=1,
                external_id=201,
                name_jp="Stadium JP",
                name_ko="Stadium KO",
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            StadiumCourseORM.__table__.insert().values(
                id=1,
                stadium_id=1,
                external_id=301,
                surface="turf",
                distance=1600,
                direction="right",
                layout="standard",
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            MatchORM.__table__.insert(),
            [
                {
                    "id": 1,
                    "name": "Split 1 Match",
                    "description": None,
                    "source_kind": "native_v2",
                    "grade": "G1",
                    "stadium_course_id": 1,
                    "scheduled_at": datetime(2026, 6, 30, 14, 59, 59),
                    "status": "scheduled",
                    "terminal_reason": None,
                    "finish_time_ms": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 2,
                    "name": "Imported Split 2 Match",
                    "description": None,
                    "source_kind": "imported_v1",
                    "grade": "LISTED",
                    "stadium_course_id": 1,
                    "scheduled_at": datetime(2026, 6, 30, 15),
                    "status": "result_confirmed",
                    "terminal_reason": None,
                    "finish_time_ms": 100000,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 3,
                    "name": "Cancelled Native Match",
                    "description": None,
                    "source_kind": "native_v2",
                    "grade": "OP",
                    "stadium_course_id": 1,
                    "scheduled_at": datetime(2026, 7, 1),
                    "status": "cancelled",
                    "terminal_reason": None,
                    "finish_time_ms": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 4,
                    "name": "Settled Native Match",
                    "description": None,
                    "source_kind": "native_v2",
                    "grade": "OP",
                    "stadium_course_id": 1,
                    "scheduled_at": datetime(2026, 7, 2),
                    "status": "settled",
                    "terminal_reason": None,
                    "finish_time_ms": 95000,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        connection.execute(
            MatchEntryORM.__table__.insert(),
            [
                {
                    "id": 1,
                    "match_id": 2,
                    "game_account_id": 1,
                    "owner_at_event_persona_id": "persona-a",
                    "affiliation_at_event": "Historical Circle",
                    "umamusume_id": 1,
                    "umamusume_variant_id": None,
                    "entry_number": 7,
                    "running_style": None,
                    "training_grade": None,
                    "rank": 1,
                    "popularity_rank": None,
                    "margin": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 2,
                    "match_id": 3,
                    "game_account_id": 1,
                    "owner_at_event_persona_id": "persona-a",
                    "affiliation_at_event": "Circle",
                    "umamusume_id": 1,
                    "umamusume_variant_id": None,
                    "entry_number": 1,
                    "running_style": None,
                    "training_grade": None,
                    "rank": None,
                    "popularity_rank": None,
                    "margin": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 3,
                    "match_id": 4,
                    "game_account_id": 1,
                    "owner_at_event_persona_id": "persona-a",
                    "affiliation_at_event": "Circle",
                    "umamusume_id": 1,
                    "umamusume_variant_id": None,
                    "entry_number": 1,
                    "running_style": None,
                    "training_grade": None,
                    "rank": 1,
                    "popularity_rank": None,
                    "margin": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        connection.execute(
            BetORM.__table__.insert(),
            [
                {
                    "id": 1,
                    "match_id": 3,
                    "persona_id": "persona-a",
                    "type": "win",
                    "selections": [2],
                    "selection_fingerprint": "b" * 64,
                    "amount": 10,
                    "status": "cancelled",
                    "active_marker": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 2,
                    "match_id": 4,
                    "persona_id": "persona-a",
                    "type": "win",
                    "selections": [3],
                    "selection_fingerprint": "d" * 64,
                    "amount": 10,
                    "status": "settled",
                    "active_marker": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        cancelled = CancelledMatch(
            match_id=3,
            match_name="Cancelled Native Match",
            previous_status=MatchStatus.BETTING_OPEN,
            status=MatchStatus.CANCELLED,
            reason=None,
            cancelled_at=NOW.replace(tzinfo=UTC),
            entry_count=1,
            cancelled_bet_count=1,
            refund_total=10,
            refunds=(
                MatchCancellationRefund(
                    persona_id="persona-a",
                    bet_ids=(1,),
                    amount=10,
                    balance_before=0,
                    balance_after=10,
                    point_transaction_id=99,
                ),
            ),
            publication=StoredMatchRefundPublication(
                publication_id=88,
                event_key="match-refund:3",
                payload_fingerprint="c" * 64,
                status=PublicationStatus.PENDING,
                target_channel_id=None,
            ),
        )
        connection.execute(
            OperationORM.__table__.insert().values(
                id=1,
                guild_id=None,
                correlation_id=None,
                actor_discord_user_id=None,
                idempotency_key=None,
                request_fingerprint=None,
                reason=None,
                created_at=NOW,
            )
        )
        connection.execute(
            MatchOperationORM.__table__.insert().values(
                operation_id=1,
                match_id=3,
                type="match_cancelled",
                before_data=None,
                after_data=cancelled.to_audit_payload(),
            )
        )
        settled = SettledMatch(
            match_id=4,
            match_name="Settled Native Match",
            previous_status=MatchStatus.RESULT_CONFIRMED,
            status=MatchStatus.SETTLED,
            grade="OP",
            settled_at=NOW.replace(tzinfo=UTC),
            result=MatchSettlementResultAuthority(1, 1, "e" * 64),
            settlement_fingerprint="f" * 64,
            active_bet_ids=(2,),
            active_stake_total=10,
            applied_odds=(
                MatchSettlementAppliedOdds(
                    bet_type=BetType.WIN,
                    selection_entry_ids=(3,),
                    selection_entry_numbers=(1,),
                    provisional_odds=Decimal("2.0000"),
                    confirmed_odds=Decimal("2.0"),
                ),
            ),
            payouts=(
                MatchSettlementPayout(
                    persona_id="persona-a",
                    bet_ids=(2,),
                    amount=20,
                    point_transaction_id=100,
                ),
            ),
            rewards=(),
            rating_rule_version=None,
            ratings=(
                MatchSettlementRating(
                    match_entry_id=3,
                    entry_number=1,
                    game_account_id=1,
                    game_account_name="Account",
                    horse_name="Horse KO",
                    affiliation_at_event="Circle",
                    rank=1,
                    rating_before=Decimal("100"),
                    base_delta=Decimal("0"),
                    adjustment_delta=Decimal("0"),
                    amount=Decimal("0"),
                    rating_after=Decimal("100"),
                    rating_transaction_id=None,
                ),
            ),
        )
        connection.execute(
            OperationORM.__table__.insert().values(
                id=2,
                guild_id=None,
                correlation_id=None,
                actor_discord_user_id=None,
                idempotency_key=None,
                request_fingerprint=None,
                reason=None,
                created_at=NOW,
            )
        )
        connection.execute(
            MatchOperationORM.__table__.insert().values(
                operation_id=2,
                match_id=4,
                type="match_settled",
                before_data=None,
                after_data=settled.to_audit_payload(),
            )
        )


def test_composed_match_export_derives_actual_kst_seasons_and_preserves_imported_raw_rows() -> None:
    runtime = DatabaseRuntime.from_engine(create_engine("sqlite+pysqlite:///:memory:"))
    Base.metadata.create_all(runtime.engine)
    try:
        _seed(runtime)
        with runtime.engine.connect() as connection:
            before = connection.scalar(select(func.count()).select_from(MatchORM))

        exports = compose_match_season_exports(runtime)
        assert [choice.key for choice in exports.search_seasons()] == ["2026-split-2", "2026-split-1"]
        artifact = exports.export_season(season_key="2026-split-2")

        with runtime.engine.connect() as connection:
            after = connection.scalar(select(func.count()).select_from(MatchORM))
        assert before == after == 4
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            match_rows = list(workbook["경기"].iter_rows(min_row=2, values_only=True))
            entry_rows = list(workbook["출전 및 결과"].iter_rows(min_row=2, values_only=True))
            assert len(match_rows) == 3
            assert match_rows[0][1] == "Imported Split 2 Match"
            assert match_rows[0][3] == "imported_v1"
            assert entry_rows[0][6:11] == (7, 1, "Account", "KR", "Owner")
            bet_rows = list(workbook["베팅 및 정산"].iter_rows(min_row=2, values_only=True))
            assert bet_rows[0][16:19] == (10, None, "2026-08-29 09:00:00 KST")
            assert bet_rows[1][13:17] == ("2.0", 20, None, None)
            assert workbook["레이팅"].max_row == 1
        finally:
            workbook.close()
    finally:
        runtime.dispose()
