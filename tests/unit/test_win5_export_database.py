"""SQLAlchemy WIN5 export projection tests on a disposable local schema."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select

from uma_st2.compose import compose_win5_season_exports
from uma_st2.infrastructure.database import Base, DatabaseRuntime
from uma_st2.infrastructure.database.orm import (
    PersonaORM,
    Win5RaceORM,
    Win5ResultORM,
    Win5RoundORM,
    Win5ScoreEventItemORM,
    Win5ScoreEventORM,
    Win5ScoreORM,
    Win5SeasonORM,
    Win5SubmissionORM,
    Win5SubmissionPickORM,
)

NOW = datetime(2026, 8, 29, 0, 0)


def _seed(runtime: DatabaseRuntime) -> None:
    with runtime.engine.begin() as connection:
        connection.execute(
            PersonaORM.__table__.insert(),
            [
                {
                    "id": "persona-a",
                    "display_name": "Alpha",
                    "status": "normal",
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": "persona-b",
                    "display_name": "Beta",
                    "status": "normal",
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        connection.execute(
            Win5SeasonORM.__table__.insert().values(
                id=1,
                name="Export Season",
                status="active",
                active_marker=True,
                starts_at=NOW,
                ends_at=None,
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            Win5RoundORM.__table__.insert().values(
                id=10,
                season_id=1,
                source_kind="imported_v1",
                type="special",
                status="scored",
                name="Mixed Void",
                opens_at=NOW,
                closes_at=NOW,
                created_at=NOW,
                updated_at=NOW,
            )
        )
        connection.execute(
            Win5RaceORM.__table__.insert(),
            [
                {
                    "id": 11,
                    "round_id": 10,
                    "name": "Void Race",
                    "scheduled_at": NOW,
                    "void_reason": "weather",
                    "voided_at": NOW,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 12,
                    "round_id": 10,
                    "name": "Live Race",
                    "scheduled_at": NOW,
                    "void_reason": None,
                    "voided_at": None,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        connection.execute(
            Win5ResultORM.__table__.insert().values(
                id=20,
                race_id=12,
                race_entry_id=None,
                gate_number=7,
                position=1,
                created_at=NOW,
            )
        )
        connection.execute(
            Win5SubmissionORM.__table__.insert(),
            [
                {
                    "id": 30,
                    "round_id": 10,
                    "persona_id": "persona-a",
                    "tier": "SPECIAL_WINNER",
                    "status": "accepted",
                    "active_marker": True,
                    "version": 1,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
                {
                    "id": 31,
                    "round_id": 10,
                    "persona_id": "persona-b",
                    "tier": "SPECIAL_WINNER",
                    "status": "cancelled",
                    "active_marker": None,
                    "version": 2,
                    "created_at": NOW,
                    "updated_at": NOW,
                },
            ],
        )
        connection.execute(
            Win5SubmissionPickORM.__table__.insert(),
            [
                {
                    "id": 40,
                    "submission_id": 30,
                    "race_id": 11,
                    "race_entry_id": None,
                    "gate_number": 3,
                    "position": 1,
                },
                {
                    "id": 41,
                    "submission_id": 30,
                    "race_id": 12,
                    "race_entry_id": None,
                    "gate_number": 7,
                    "position": 1,
                },
                {
                    "id": 42,
                    "submission_id": 31,
                    "race_id": 11,
                    "race_entry_id": None,
                    "gate_number": 9,
                    "position": 1,
                },
            ],
        )
        connection.execute(
            Win5ScoreEventORM.__table__.insert().values(
                id=50,
                operation_id=None,
                season_id=1,
                round_id=10,
                race_id=None,
                submission_id=30,
                submission_version=1,
                persona_id="persona-a",
                tier="SPECIAL_WINNER",
                result_fingerprint="a" * 64,
                scoring_policy_version="special/v1",
                reward_policy_version="reward/v1",
                exact_count=1,
                wrong_position_count=0,
                off_board_count=0,
                missing_count=0,
                season_score_delta=1,
                top1_score_delta=1,
                circle_point_reward=0,
                created_at=NOW,
            )
        )
        connection.execute(
            Win5ScoreEventItemORM.__table__.insert(),
            [
                {
                    "id": 60,
                    "score_event_id": 50,
                    "race_id": 11,
                    "position": 1,
                    "submission_pick_id": 40,
                    "matched_result_id": None,
                    "outcome": "void",
                    "season_score_delta": 0,
                },
                {
                    "id": 61,
                    "score_event_id": 50,
                    "race_id": 12,
                    "position": 1,
                    "submission_pick_id": 41,
                    "matched_result_id": 20,
                    "outcome": "exact",
                    "season_score_delta": 1,
                },
            ],
        )
        connection.execute(
            Win5ScoreORM.__table__.insert().values(
                season_id=1,
                persona_id="persona-a",
                season_score=1,
                top1_score=1,
                updated_at=NOW,
            )
        )


def test_composed_export_materializes_complete_detached_workbook_without_writes() -> None:
    runtime = DatabaseRuntime.from_engine(create_engine("sqlite+pysqlite:///:memory:"))
    Base.metadata.create_all(runtime.engine)
    try:
        _seed(runtime)
        with runtime.engine.connect() as connection:
            before = connection.scalar(select(func.count()).select_from(Win5SubmissionORM))

        exports = compose_win5_season_exports(runtime)
        assert [choice.id for choice in exports.search_seasons()] == [1]
        artifact = exports.export_season(season_id=1)

        with runtime.engine.connect() as connection:
            after = connection.scalar(select(func.count()).select_from(Win5SubmissionORM))
        assert before == after == 2
        assert artifact.scope_id == "1"
        workbook = load_workbook(BytesIO(artifact.content), read_only=True)
        try:
            judgement_rows = list(workbook["제출 및 판정"].iter_rows(min_row=2, values_only=True))
            assert len(judgement_rows) == 3
            assert any(row[16] == "void" and row[15] == "void" for row in judgement_rows)
            assert any(row[16] == "exact" and row[15] == "7" for row in judgement_rows)
            assert any(row[5] == "cancelled" and row[16] is None for row in judgement_rows)
        finally:
            workbook.close()
    finally:
        runtime.dispose()
