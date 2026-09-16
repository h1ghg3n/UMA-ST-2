"""Circle Point current export Application and workbook v1 tests."""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import load_workbook

from uma_st2.application.execution import QueryRunner
from uma_st2.application.exporting import (
    CIRCLE_POINT_EXPORT_PROJECTION_VERSION,
    CIRCLE_POINT_EXPORT_SCOPE_ID,
    CIRCLE_POINT_EXPORT_SCOPE_NAME,
    CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION,
    CirclePointExportInvalidSourceError,
    CirclePointExportProjection,
    CirclePointExports,
    CirclePointExportSource,
    CirclePointExportTransaction,
    CirclePointExportWallet,
    ExportArtifact,
)
from uma_st2.domain.identity import PersonaStatus
from uma_st2.infrastructure.exporting import (
    CirclePointXlsxRenderer,
    CirclePointXlsxRenderError,
    circle_point_xlsx,
)

SOURCE_CUTOFF = datetime(2026, 8, 29, 1, 2, 3, tzinfo=UTC)
GENERATED_AT = datetime(2026, 8, 29, 1, 3, 4, tzinfo=UTC)


def _source() -> CirclePointExportSource:
    return CirclePointExportSource(
        source_cutoff=SOURCE_CUTOFF,
        wallets=(
            CirclePointExportWallet(
                persona_id="persona-a",
                display_name="=Formula Owner",
                status=PersonaStatus.NORMAL,
                balance=0,
                updated_at=datetime(2026, 8, 28, 23, tzinfo=UTC),
            ),
            CirclePointExportWallet(
                persona_id="persona-b",
                display_name="Owner B",
                status=PersonaStatus.WARNING,
                balance=100,
                updated_at=datetime(2026, 8, 29, tzinfo=UTC),
            ),
        ),
        transactions=(
            CirclePointExportTransaction(
                id=10,
                persona_id="persona-a",
                display_name="=Formula Owner",
                action="manual_grant",
                amount=50,
                created_at=datetime(2026, 8, 28, 23, 30, tzinfo=UTC),
                operation_id=20,
                reason="@retained reason",
            ),
            CirclePointExportTransaction(
                id=11,
                persona_id="persona-b",
                display_name="Owner B",
                action="match_bet_stake",
                amount=-20,
                created_at=datetime(2026, 8, 29, 0, 30, tzinfo=UTC),
            ),
        ),
    )


class FakeRepository:
    def __init__(self, source: CirclePointExportSource) -> None:
        self.source = source

    def get_current_source(self, *, source_cutoff: datetime) -> CirclePointExportSource:
        assert source_cutoff == SOURCE_CUTOFF
        return self.source


class FakeUnitOfWork:
    def __init__(self, source: CirclePointExportSource) -> None:
        self.repository = FakeRepository(source)
        self.active = False
        self.rollback_calls = 0

    @property
    def circle_point_exports(self) -> FakeRepository:
        assert self.active
        return self.repository

    def __enter__(self):  # type: ignore[no-untyped-def]
        self.active = True
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> bool:  # type: ignore[no-untyped-def]
        self.active = False
        return False

    def commit(self) -> None:
        raise AssertionError("Export query must not commit.")

    def rollback(self) -> None:
        self.rollback_calls += 1


@dataclass
class RecordingRenderer:
    unit_of_work: FakeUnitOfWork
    projection: CirclePointExportProjection | None = None

    def render(
        self,
        projection: CirclePointExportProjection,
        *,
        generated_at: datetime,
    ) -> ExportArtifact:
        assert not self.unit_of_work.active
        self.projection = projection
        content = b"circle-point-workbook"
        return ExportArtifact(
            filename="circle-points.xlsx",
            media_type="application/xlsx",
            content=content,
            sha256_hex=sha256(content).hexdigest(),
            generated_at=generated_at,
            source_cutoff=projection.source_cutoff,
            schema_version=CIRCLE_POINT_EXPORT_WORKBOOK_SCHEMA_VERSION,
            projection_version=CIRCLE_POINT_EXPORT_PROJECTION_VERSION,
            scope_type="circle_points",
            scope_id=CIRCLE_POINT_EXPORT_SCOPE_ID,
            scope_name=CIRCLE_POINT_EXPORT_SCOPE_NAME,
            row_count=projection.workbook_data_row_count,
        )


def _exports(
    source: CirclePointExportSource,
) -> tuple[CirclePointExports, FakeUnitOfWork, RecordingRenderer]:
    unit_of_work = FakeUnitOfWork(source)
    renderer = RecordingRenderer(unit_of_work)
    clock_values = iter((SOURCE_CUTOFF, GENERATED_AT))
    return (
        CirclePointExports(
            QueryRunner(lambda: unit_of_work),  # type: ignore[arg-type]
            renderer,
            clock=lambda: next(clock_values),
        ),
        unit_of_work,
        renderer,
    )


def test_application_closes_query_uow_and_does_not_reconcile_retained_sum() -> None:
    exports, unit_of_work, renderer = _exports(_source())

    artifact = exports.export_current()

    assert unit_of_work.rollback_calls == 1
    assert renderer.projection is not None
    assert renderer.projection.current_balance_total == 100
    assert sum(transaction.amount for transaction in renderer.projection.transactions) == 30
    assert artifact.row_count == 4


def test_application_preserves_signed_authoritative_current_balance() -> None:
    source = _source()
    signed_source = replace(
        source,
        wallets=(replace(source.wallets[0], balance=-5), source.wallets[1]),
    )
    exports, _, renderer = _exports(signed_source)

    exports.export_current()

    assert renderer.projection is not None
    assert renderer.projection.current_balance_total == 95
    assert renderer.projection.wallets[0].balance == -5


def test_application_rejects_retained_transaction_without_current_wallet() -> None:
    source = _source()
    missing_owner = replace(
        source.transactions[0],
        persona_id="persona-missing",
        display_name="Missing",
    )
    exports, unit_of_work, renderer = _exports(replace(source, transactions=(missing_owner,)))

    with pytest.raises(CirclePointExportInvalidSourceError):
        exports.export_current()

    assert unit_of_work.rollback_calls == 1
    assert renderer.projection is None


def test_xlsx_v1_preserves_current_balances_and_only_retained_history() -> None:
    exports, _, recording_renderer = _exports(_source())
    exports.export_current()
    projection = recording_renderer.projection
    assert projection is not None

    artifact = CirclePointXlsxRenderer().render(projection, generated_at=GENERATED_AT)

    assert artifact.filename == "circle-points_20260829T010304Z.xlsx"
    assert artifact.sha256_hex == sha256(artifact.content).hexdigest()
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    try:
        assert workbook.sheetnames == ["요약", "현재 잔액", "거래 내역"]
        assert workbook["요약"]["B8"].value == 100
        wallet_rows = list(workbook["현재 잔액"].iter_rows(min_row=2, values_only=True))
        assert wallet_rows[0] == (
            "persona-a",
            "'=Formula Owner",
            "normal",
            0,
            "2026-08-29 08:00:00 KST",
        )
        transaction_rows = list(workbook["거래 내역"].iter_rows(min_row=2, values_only=True))
        assert transaction_rows[0][3:] == (
            "manual_grant",
            50,
            "2026-08-29 08:30:00 KST",
            20,
            "'@retained reason",
        )
        assert transaction_rows[1][3:] == (
            "match_bet_stake",
            -20,
            "2026-08-29 09:30:00 KST",
            None,
            None,
        )
    finally:
        workbook.close()


def test_xlsx_v1_rejects_rows_that_cannot_fit_without_truncation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    exports, _, recording_renderer = _exports(_source())
    exports.export_current()
    projection = recording_renderer.projection
    assert projection is not None
    monkeypatch.setattr(circle_point_xlsx, "_EXCEL_MAX_ROWS", 2)

    with pytest.raises(CirclePointXlsxRenderError):
        CirclePointXlsxRenderer().render(projection, generated_at=GENERATED_AT)


def test_empty_current_snapshot_renders_header_only_tables() -> None:
    empty = CirclePointExportSource(source_cutoff=SOURCE_CUTOFF)
    unit_of_work = FakeUnitOfWork(empty)
    clock_values = iter((SOURCE_CUTOFF, GENERATED_AT))
    exports = CirclePointExports(
        QueryRunner(lambda: unit_of_work),  # type: ignore[arg-type]
        CirclePointXlsxRenderer(),
        clock=lambda: next(clock_values),
    )

    artifact = exports.export_current()

    assert artifact.row_count == 0
    workbook = load_workbook(BytesIO(artifact.content), read_only=True)
    try:
        assert workbook["현재 잔액"].max_row == 1
        assert workbook["거래 내역"].max_row == 1
    finally:
        workbook.close()
