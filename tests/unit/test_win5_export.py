"""WIN5 Season export application and workbook v1 tests."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from hashlib import sha256
from io import BytesIO

import pytest
from openpyxl import load_workbook

from uma_st2.application.execution import QueryRunner
from uma_st2.application.exporting import (
    WIN5_EXPORT_PROJECTION_VERSION,
    WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION,
    ExportArtifact,
    Win5ExportEntry,
    Win5ExportJudgementItem,
    Win5ExportPick,
    Win5ExportRace,
    Win5ExportResult,
    Win5ExportRound,
    Win5ExportScore,
    Win5ExportScoreEvent,
    Win5ExportSeasonChoice,
    Win5ExportSubmission,
    Win5SeasonExportInvalidSourceError,
    Win5SeasonExports,
    Win5SeasonExportSource,
    Win5SeasonExportUnavailableError,
)
from uma_st2.domain.win5 import (
    Win5JudgementOutcome,
    Win5RoundStatus,
    Win5RoundType,
    Win5SeasonStatus,
    Win5SubmissionStatus,
    Win5SubmissionTier,
)
from uma_st2.infrastructure.exporting import Win5SeasonXlsxRenderer

NOW = datetime(2026, 8, 29, 1, 2, 3, tzinfo=UTC)
GENERATED_AT = datetime(2026, 8, 29, 1, 3, 4, tzinfo=UTC)


def _normal_round() -> Win5ExportRound:
    entries = tuple(
        Win5ExportEntry(
            id=100 + position,
            gate_number=position,
            name="=formula horse" if position == 1 else f"Horse {position}",
        )
        for position in range(1, 6)
    )
    return Win5ExportRound(
        id=10,
        round_type=Win5RoundType.NORMAL,
        status=Win5RoundStatus.SCORED,
        name="Normal Final",
        opens_at=NOW,
        closes_at=NOW,
        races=(
            Win5ExportRace(
                id=11,
                name="Normal Race",
                scheduled_at=NOW,
                void_reason=None,
                voided_at=None,
                entries=entries,
                results=tuple(
                    Win5ExportResult(
                        id=200 + position,
                        race_id=11,
                        position=position,
                        race_entry_id=100 + position,
                        gate_number=None,
                    )
                    for position in range(1, 6)
                ),
            ),
        ),
    )


def _special_round() -> Win5ExportRound:
    return Win5ExportRound(
        id=20,
        round_type=Win5RoundType.SPECIAL,
        status=Win5RoundStatus.SCORED,
        name="Mixed Void Special",
        opens_at=NOW,
        closes_at=NOW,
        races=(
            Win5ExportRace(
                id=21,
                name="Race A",
                scheduled_at=NOW,
                void_reason="official no-contest",
                voided_at=NOW,
            ),
            Win5ExportRace(
                id=22,
                name="Race B",
                scheduled_at=NOW,
                void_reason=None,
                voided_at=None,
                results=(
                    Win5ExportResult(
                        id=221,
                        race_id=22,
                        position=1,
                        race_entry_id=None,
                        gate_number=7,
                    ),
                ),
            ),
        ),
    )


def _cancelled_round() -> Win5ExportRound:
    return Win5ExportRound(
        id=30,
        round_type=Win5RoundType.SPECIAL,
        status=Win5RoundStatus.CANCELLED,
        name="All Void Special",
        opens_at=NOW,
        closes_at=NOW,
        races=(
            Win5ExportRace(
                id=31,
                name="Cancelled Race",
                scheduled_at=NOW,
                void_reason="weather",
                voided_at=NOW,
            ),
        ),
    )


def _normal_submission() -> Win5ExportSubmission:
    picks = tuple(
        Win5ExportPick(
            id=300 + position,
            race_id=11,
            position=position,
            race_entry_id=100 + position,
            gate_number=None,
        )
        for position in range(1, 6)
    )
    return Win5ExportSubmission(
        id=40,
        round_id=10,
        persona_id="persona-a",
        display_name="=formula member",
        tier=Win5SubmissionTier.TOP5,
        status=Win5SubmissionStatus.ACCEPTED,
        version=2,
        created_at=NOW,
        updated_at=NOW,
        picks=picks,
        score_event=Win5ExportScoreEvent(
            id=50,
            season_id=1,
            round_id=10,
            persona_id="persona-a",
            submission_version=2,
            tier=Win5SubmissionTier.TOP5,
            exact_count=5,
            wrong_position_count=0,
            off_board_count=0,
            missing_count=0,
            season_score_delta=15,
            top1_score_delta=1,
            circle_point_reward=50,
            created_at=NOW,
            items=tuple(
                Win5ExportJudgementItem(
                    id=400 + position,
                    race_id=11,
                    position=position,
                    submission_pick_id=300 + position,
                    matched_result_id=200 + position,
                    outcome=Win5JudgementOutcome.EXACT,
                    season_score_delta=position,
                )
                for position in range(1, 6)
            ),
        ),
    )


def _special_submission() -> Win5ExportSubmission:
    return Win5ExportSubmission(
        id=41,
        round_id=20,
        persona_id="persona-b",
        display_name="Special Member",
        tier=Win5SubmissionTier.SPECIAL_WINNER,
        status=Win5SubmissionStatus.ACCEPTED,
        version=1,
        created_at=NOW,
        updated_at=NOW,
        picks=(
            Win5ExportPick(id=411, race_id=21, position=1, race_entry_id=None, gate_number=3),
            Win5ExportPick(id=412, race_id=22, position=1, race_entry_id=None, gate_number=7),
        ),
        score_event=Win5ExportScoreEvent(
            id=51,
            season_id=1,
            round_id=20,
            persona_id="persona-b",
            submission_version=1,
            tier=Win5SubmissionTier.SPECIAL_WINNER,
            exact_count=1,
            wrong_position_count=0,
            off_board_count=0,
            missing_count=0,
            season_score_delta=1,
            top1_score_delta=1,
            circle_point_reward=0,
            created_at=NOW,
            items=(
                Win5ExportJudgementItem(
                    id=511,
                    race_id=21,
                    position=1,
                    submission_pick_id=411,
                    matched_result_id=None,
                    outcome=Win5JudgementOutcome.VOID,
                    season_score_delta=0,
                ),
                Win5ExportJudgementItem(
                    id=512,
                    race_id=22,
                    position=1,
                    submission_pick_id=412,
                    matched_result_id=221,
                    outcome=Win5JudgementOutcome.EXACT,
                    season_score_delta=1,
                ),
            ),
        ),
    )


def _missing_submission() -> Win5ExportSubmission:
    return Win5ExportSubmission(
        id=42,
        round_id=20,
        persona_id="persona-c",
        display_name="Missing Member",
        tier=Win5SubmissionTier.SPECIAL_WINNER,
        status=Win5SubmissionStatus.ACCEPTED,
        version=1,
        created_at=NOW,
        updated_at=NOW,
        picks=(Win5ExportPick(id=421, race_id=21, position=1, race_entry_id=None, gate_number=9),),
        score_event=Win5ExportScoreEvent(
            id=52,
            season_id=1,
            round_id=20,
            persona_id="persona-c",
            submission_version=1,
            tier=Win5SubmissionTier.SPECIAL_WINNER,
            exact_count=0,
            wrong_position_count=0,
            off_board_count=0,
            missing_count=1,
            season_score_delta=0,
            top1_score_delta=0,
            circle_point_reward=0,
            created_at=NOW,
            items=(
                Win5ExportJudgementItem(
                    id=521,
                    race_id=21,
                    position=1,
                    submission_pick_id=421,
                    matched_result_id=None,
                    outcome=Win5JudgementOutcome.VOID,
                    season_score_delta=0,
                ),
                Win5ExportJudgementItem(
                    id=522,
                    race_id=22,
                    position=1,
                    submission_pick_id=None,
                    matched_result_id=None,
                    outcome=Win5JudgementOutcome.MISSING,
                    season_score_delta=0,
                ),
            ),
        ),
    )


def _cancelled_submission() -> Win5ExportSubmission:
    return Win5ExportSubmission(
        id=43,
        round_id=10,
        persona_id="persona-d",
        display_name="Cancelled Member",
        tier=Win5SubmissionTier.TOP1,
        status=Win5SubmissionStatus.CANCELLED,
        version=3,
        created_at=NOW,
        updated_at=NOW,
        picks=(
            Win5ExportPick(
                id=431,
                race_id=11,
                position=1,
                race_entry_id=101,
                gate_number=None,
            ),
        ),
    )


def _source() -> Win5SeasonExportSource:
    return Win5SeasonExportSource(
        season_id=1,
        season_name="2026 / Summer",
        season_status=Win5SeasonStatus.ACTIVE,
        starts_at=NOW,
        ends_at=None,
        source_cutoff=NOW,
        rounds=(_normal_round(), _special_round(), _cancelled_round()),
        submissions=(
            _normal_submission(),
            _special_submission(),
            _missing_submission(),
            _cancelled_submission(),
        ),
        scores=(
            Win5ExportScore("persona-a", "=formula member", 15, 1),
            Win5ExportScore("persona-b", "Special Member", 1, 1),
            Win5ExportScore("persona-c", "Missing Member", 1, 0),
        ),
    )


class FakeRepository:
    def __init__(self, source: Win5SeasonExportSource | None) -> None:
        self.source = source

    def search_seasons(self, *, query: str, limit: int) -> tuple[Win5ExportSeasonChoice, ...]:
        choices = (
            Win5ExportSeasonChoice(1, "Active", Win5SeasonStatus.ACTIVE),
            Win5ExportSeasonChoice(2, "Closed", Win5SeasonStatus.CLOSED),
        )
        return tuple(choice for choice in choices if query.lower() in choice.name.lower())[:limit]

    def get_season_source(
        self,
        *,
        season_id: int,
        source_cutoff: datetime,
    ) -> Win5SeasonExportSource | None:
        assert season_id == 1
        assert source_cutoff == NOW
        return self.source


class FakeUnitOfWork:
    def __init__(self, repository: FakeRepository) -> None:
        self.repository = repository
        self.active = False
        self.rollback_calls = 0

    @property
    def win5_season_exports(self) -> FakeRepository:
        assert self.active
        return self.repository

    def __enter__(self) -> FakeUnitOfWork:
        assert not self.active
        self.active = True
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> bool:  # type: ignore[no-untyped-def]
        self.active = False
        return False

    def commit(self) -> None:
        raise AssertionError("Export query must not commit.")

    def rollback(self) -> None:
        self.rollback_calls += 1


@dataclass
class RecordingRenderer:
    unit_of_work: FakeUnitOfWork
    projection: object | None = None

    def render(self, projection, *, generated_at: datetime) -> ExportArtifact:  # type: ignore[no-untyped-def]
        assert not self.unit_of_work.active
        self.projection = projection
        content = b"workbook"
        return ExportArtifact(
            filename="win5.xlsx",
            media_type="application/xlsx",
            content=content,
            sha256_hex=sha256(content).hexdigest(),
            generated_at=generated_at,
            source_cutoff=projection.source_cutoff,
            schema_version=WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION,
            projection_version=WIN5_EXPORT_PROJECTION_VERSION,
            scope_type="win5_season",
            scope_id=str(projection.season_id),
            scope_name=projection.season_name,
            row_count=projection.workbook_data_row_count,
        )


def _exports(source: Win5SeasonExportSource | None) -> tuple[Win5SeasonExports, FakeUnitOfWork, RecordingRenderer]:
    unit_of_work = FakeUnitOfWork(FakeRepository(source))
    renderer = RecordingRenderer(unit_of_work)
    clock_values = iter((NOW, GENERATED_AT))
    exports = Win5SeasonExports(
        QueryRunner(lambda: unit_of_work),  # type: ignore[arg-type]
        renderer,
        clock=lambda: next(clock_values),
    )
    return exports, unit_of_work, renderer


def test_application_closes_query_uow_before_render_and_derives_persisted_rankings() -> None:
    exports, unit_of_work, renderer = _exports(_source())

    artifact = exports.export_season(season_id=1)

    projection = renderer.projection
    assert projection is not None
    assert artifact.generated_at == GENERATED_AT
    assert unit_of_work.rollback_calls == 1
    assert [(entry.rank, entry.persona_id, entry.score) for entry in projection.season_standings] == [
        (1, "persona-a", 15),
        (2, "persona-b", 1),
        (2, "persona-c", 1),
    ]
    assert projection.hall_of_fame_submission_ids == (40,)
    assert projection.participant_count == 4
    assert projection.judgement_row_count == 10


def test_application_searches_only_valid_export_choices() -> None:
    exports, unit_of_work, _ = _exports(_source())

    assert exports.search_seasons(query="closed") == (Win5ExportSeasonChoice(2, "Closed", Win5SeasonStatus.CLOSED),)
    assert unit_of_work.rollback_calls == 1


def test_application_rejects_unavailable_or_incomplete_source_before_render() -> None:
    exports, _, renderer = _exports(None)
    with pytest.raises(Win5SeasonExportUnavailableError):
        exports.export_season(season_id=1)
    assert renderer.projection is None

    malformed = _source()
    malformed = Win5SeasonExportSource(
        season_id=malformed.season_id,
        season_name=malformed.season_name,
        season_status=malformed.season_status,
        starts_at=malformed.starts_at,
        ends_at=malformed.ends_at,
        source_cutoff=malformed.source_cutoff,
        rounds=malformed.rounds,
        submissions=(
            Win5ExportSubmission(
                id=99,
                round_id=10,
                persona_id="broken",
                display_name="Broken",
                tier=Win5SubmissionTier.TOP1,
                status=Win5SubmissionStatus.ACCEPTED,
                version=1,
                created_at=NOW,
                updated_at=NOW,
            ),
        ),
        scores=(),
    )
    exports, _, renderer = _exports(malformed)
    with pytest.raises(Win5SeasonExportInvalidSourceError):
        exports.export_season(season_id=1)
    assert renderer.projection is None


def test_xlsx_v1_preserves_void_missing_cancelled_and_formula_safe_values() -> None:
    exports, _, recording_renderer = _exports(_source())
    exports.export_season(season_id=1)
    projection = recording_renderer.projection
    assert projection is not None

    artifact = Win5SeasonXlsxRenderer().render(projection, generated_at=GENERATED_AT)

    assert artifact.filename == "win5_1_2026_Summer_20260829T010304Z.xlsx"
    assert artifact.schema_version == WIN5_EXPORT_WORKBOOK_SCHEMA_VERSION
    assert artifact.sha256_hex == sha256(artifact.content).hexdigest()
    workbook = load_workbook(BytesIO(artifact.content), data_only=False)
    try:
        assert workbook.sheetnames == ["시즌 요약", "라운드", "제출 및 판정", "명예의 전당"]
        summary = workbook["시즌 요약"]
        assert summary["B3"].value == "2026 / Summer"
        assert summary["B7"].value == "2026-08-29 10:03:04 KST"
        assert summary["B8"].value == "2026-08-29 10:02:03 KST"
        assert summary["B18"].value == "'=formula member"
        assert len(summary._charts) == 1

        rounds = workbook["라운드"]
        round_rows = list(rounds.iter_rows(min_row=2, values_only=True))
        assert any(row[8] == "Race A" and row[10] == "Y" and row[13] == "void" for row in round_rows)
        assert any(row[8] == "Cancelled Race" and row[3] == "cancelled" and row[11] == "weather" for row in round_rows)

        judgement = workbook["제출 및 판정"]
        rows = list(judgement.iter_rows(min_row=2, values_only=True))
        assert len(rows) == projection.judgement_row_count
        assert any(row[16] == "void" and row[14] == "3" and row[15] == "void" for row in rows)
        assert any(row[16] == "missing" and row[14] is None and row[12] == "Race B" for row in rows)
        assert any(row[5] == "cancelled" and row[16] is None for row in rows)
        assert any(row[14] == "1 · =formula horse" for row in rows)

        hall = workbook["명예의 전당"]
        assert hall.max_row == 2
        assert hall["D2"].value == "'=formula member"
    finally:
        workbook.close()
